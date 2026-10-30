"""Gera DOC/GUIA_CRIACAO_DE_TESTES.docx — versao didatica completa."""

import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Margens ──────────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3.0)
    section.right_margin  = Cm(2.5)

# ── Estilo base ───────────────────────────────────────────────────────────────
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)

# Estilos usados pelo Word ao montar o sumário (capítulos em negrito, seções recuadas)
from docx.enum.style import WD_STYLE_TYPE  # noqa: E402

for _nome, _recuo, _negrito in (('toc 1', 0.0, True), ('toc 2', 0.6, False)):
    try:
        _st = doc.styles[_nome]
    except KeyError:
        _st = doc.styles.add_style(_nome, WD_STYLE_TYPE.PARAGRAPH)
        _st.base_style = doc.styles['Normal']
    _st.font.bold = _negrito
    _st.paragraph_format.left_indent = Cm(_recuo)
    _st.paragraph_format.space_before = Pt(4 if _negrito else 0)
    _st.paragraph_format.space_after = Pt(2)

# Legendas: o Word reaplica o estilo ao número (campo SEQ) quando atualiza
_cap = doc.styles['Caption']
_cap.font.bold = False
_cap.font.italic = True
_cap.font.size = Pt(9)
_cap.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

# Cores
AZUL_ESCURO  = RGBColor(0x1F, 0x39, 0x64)
AZUL_MEDIO   = RGBColor(0x2E, 0x74, 0xB5)
CINZA_TEXTO  = RGBColor(0x55, 0x55, 0x55)
VERDE        = RGBColor(0x00, 0x70, 0x00)
ROXO         = RGBColor(0x80, 0x00, 0x80)
LARANJA      = RGBColor(0xC0, 0x50, 0x00)
PRETO        = RGBColor(0x00, 0x00, 0x00)
VERMELHO     = RGBColor(0xCC, 0x00, 0x00)
FUNDO_CODE   = 'F0F0F0'
FUNDO_HDR    = '1F3964'
FUNDO_ALT    = 'D6E4F0'
FUNDO_DICA   = 'FFF8DC'
FUNDO_AVISO  = 'FFF0F0'


# ══════════════════════════════════════════════════════════════════════════════
# Helpers
# ══════════════════════════════════════════════════════════════════════════════

def _shading(cell_or_para, fill_hex):
    el = cell_or_para if hasattr(cell_or_para, '_tc') else cell_or_para._p
    pPr = el.get_or_add_pPr() if hasattr(el, 'get_or_add_pPr') else el
    if hasattr(cell_or_para, '_tc'):
        pPr = cell_or_para._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    pPr.append(shd)


_QUEBRA = [False]


def quebra():
    """Próximo título começa em nova página (page_break_before, sem parágrafo vazio).

    Remove os espaçadores vazios do fim do capítulo: se transbordassem para a
    página seguinte, gerariam uma página em branco antes do próximo título.
    """
    body = doc.element.body
    while True:
        last = body[-2] if body[-1].tag == qn('w:sectPr') else body[-1]
        if last.tag != qn('w:p') or last.findall(qn('w:r')):
            break
        body.remove(last)
    _QUEBRA[0] = True


def _consome_quebra(p):
    if _QUEBRA[0]:
        p.paragraph_format.page_break_before = True
        _QUEBRA[0] = False


def h1(text):
    p = doc.add_heading('', level=1)
    p.clear()
    _consome_quebra(p)
    r = p.add_run(text)
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = AZUL_ESCURO
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after  = Pt(6)
    return p


def h2(text):
    p = doc.add_heading('', level=2)
    p.clear()
    _consome_quebra(p)
    r = p.add_run(text)
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = AZUL_MEDIO
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after  = Pt(4)
    return p


def h3(text):
    p = doc.add_heading('', level=3)
    p.clear()
    r = p.add_run(text)
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = AZUL_MEDIO
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after  = Pt(2)
    return p


def para(text, bold=False, italic=False, color=None, size=11,
         space_before=0, space_after=6, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.color.rgb = color or PRETO
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    if align:
        p.alignment = align
    return p


def mixed_para(*parts, space_before=0, space_after=6):
    """parts = list of (text, bold, italic, color)"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    for text, bold, italic, color in parts:
        r = p.add_run(text)
        r.bold   = bold
        r.italic = italic
        r.font.size = Pt(11)
        r.font.color.rgb = color or PRETO
    return p


def bullet(text, level=0, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Cm(0.5 + level * 0.5)
    p.paragraph_format.space_after = Pt(3)
    if bold_prefix:
        r = p.add_run(bold_prefix + ' ')
        r.bold = True
        r.font.size = Pt(11)
    r2 = p.add_run(text)
    r2.font.size = Pt(11)
    return p


def numbered(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(4)
    if bold_prefix:
        r = p.add_run(bold_prefix + '  ')
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = AZUL_MEDIO
    r2 = p.add_run(text)
    r2.font.size = Pt(11)
    return p


def box(text, fill=FUNDO_DICA, title=None, title_color=None):
    """Caixa de destaque (dica, aviso, nota)."""
    if title:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent  = Cm(0.8)
        p.paragraph_format.right_indent = Cm(0.8)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after  = Pt(0)
        r = p.add_run('  ' + title)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = title_color or AZUL_ESCURO
        _shading_p(p, fill)

    for line in text.split('\n'):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent  = Cm(0.8)
        p.paragraph_format.right_indent = Cm(0.8)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(0)
        r = p.add_run('  ' + line)
        r.font.size = Pt(10)
        r.font.color.rgb = PRETO
        _shading_p(p, fill)

    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_after = Pt(8)


def _shading_p(p, fill_hex):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    pPr.append(shd)


def code_block(lines_with_tokens, caption=None):
    """
    lines_with_tokens: lista de linhas.
    Cada linha pode ser string pura ou lista de (texto, cor).
    """
    if caption:
        p = doc.add_paragraph()
        r = p.add_run('  ' + caption)
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = CINZA_TEXTO
        p.paragraph_format.space_after = Pt(0)

    for line in lines_with_tokens:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent  = Cm(0.5)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(0)
        _shading_p(p, FUNDO_CODE)

        if isinstance(line, str):
            r = p.add_run(line if line else ' ')
            r.font.name  = 'Courier New'
            r.font.size  = Pt(9)
            r.font.color.rgb = PRETO
        elif not line:
            r = p.add_run(' ')
            r.font.name = 'Courier New'
            r.font.size = Pt(9)
        else:
            for item in line:
                if isinstance(item, tuple):
                    text, color = item
                else:
                    text, color = item, PRETO
                r = p.add_run(text if text else ' ')
                r.font.name  = 'Courier New'
                r.font.size  = Pt(9)
                r.font.color.rgb = color or PRETO

    doc.add_paragraph().paragraph_format.space_after = Pt(6)


# Atalho: código simples (string multiline → lista de linhas)
def code(text, caption=None):
    code_block(text.split('\n'), caption=caption)


# Tokens de sintaxe Python (coloracao manual)
KW  = lambda t: (t, RGBColor(0x00, 0x00, 0xCC))   # azul — keywords
STR = lambda t: (t, RGBColor(0xA0, 0x20, 0x00))   # vermelho-marrom — strings
CMT = lambda t: (t, RGBColor(0x00, 0x70, 0x00))   # verde — comentarios
NUM = lambda t: (t, RGBColor(0x09, 0x86, 0x58))   # verde-azul — numeros
DEF = lambda t: (t, RGBColor(0x1A, 0x5C, 0x86))   # azul — definicoes
TXT = lambda t: (t, PRETO)                         # texto normal


def table(headers, rows, widths=None, titulo=None):
    if titulo:
        pc = doc.add_paragraph(style='Caption')
        pc.paragraph_format.space_before = Pt(6)
        pc.paragraph_format.space_after  = Pt(2)
        pc.paragraph_format.keep_with_next = True
        _rotulo(pc, 'Tabela', titulo)
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = 'Table Grid'
    # Cabeçalho
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ''
        r = cell.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), FUNDO_HDR)
        cell._tc.get_or_add_tcPr().append(shd)
    # Linhas
    for ri, row in enumerate(rows):
        cells = t.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = ''
            r = cells[ci].paragraphs[0].add_run(val)
            r.font.size = Pt(10)
            if ri % 2 == 0:
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:color'), 'auto')
                shd.set(qn('w:fill'), FUNDO_ALT)
                cells[ci]._tc.get_or_add_tcPr().append(shd)
    if widths:
        for row in t.rows:
            for ci, w in enumerate(widths):
                if ci < len(row.cells):
                    row.cells[ci].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def sep():
    """Linha separadora visual."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    r = p.add_run('─' * 80)
    r.font.size  = Pt(7)
    r.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)


# Capturas geradas por DOC/_gerar_capturas.py (dados simulados)
IMG = 'DOC/img/'

# Contadores das legendas — o número fica no campo SEQ e é recalculado pelo Word
_SEQ = {'Figura': 0, 'Tabela': 0}
_PREFIXO = re.compile(r'^(Figura|Tabela) \d+ — ')


def _campo(p, instr, texto=''):
    """Campo do Word (SEQ, TOC...) com texto em cache; retorna o run visível."""
    def fld(tipo):
        el = OxmlElement('w:fldChar')
        el.set(qn('w:fldCharType'), tipo)
        p.add_run()._r.append(el)

    fld('begin')
    instr_el = OxmlElement('w:instrText')
    instr_el.set(qn('xml:space'), 'preserve')
    instr_el.text = f' {instr} '
    p.add_run()._r.append(instr_el)
    fld('separate')
    r = p.add_run(texto)
    fld('end')
    return r


def _rotulo(p, tipo, texto, size=9):
    """'Figura N — texto' / 'Tabela N — texto' com N em campo SEQ (base das listas)."""
    _SEQ[tipo] += 1
    partes = [p.add_run(f'{tipo} ')]
    partes.append(_campo(p, f'SEQ {tipo} \\* ARABIC', str(_SEQ[tipo])))
    partes.append(p.add_run(' — ' + _PREFIXO.sub('', texto)))
    for r in partes:
        r.italic = True
        r.font.size = Pt(size)
        r.font.color.rgb = CINZA_TEXTO


def _legenda(text, tipo='Figura'):
    p = doc.add_paragraph(style='Caption')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(10)
    _rotulo(p, tipo, text)


def titulo_lista(text):
    """Título com aparência de h1, mas fora do sumário (Sumário, Lista de Figuras...)."""
    p = doc.add_paragraph()
    _consome_quebra(p)
    r = p.add_run(text)
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = AZUL_ESCURO
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after  = Pt(6)
    return p


def indice(instr, aviso):
    """Sumário/lista gerada pelo Word a partir de um campo TOC."""
    p = doc.add_paragraph()
    r = _campo(p, instr, aviso)
    r.italic = True
    r.font.color.rgb = CINZA_TEXTO
    # Separa o campo do conteúdo seguinte (a atualização substitui o campo inteiro).
    # Linha de 1 pt: não transborda sozinha para uma página nova.
    sep = doc.add_paragraph()
    sep.paragraph_format.space_before = Pt(0)
    sep.paragraph_format.space_after = Pt(0)
    sep.paragraph_format.line_spacing = Pt(1)
    sep.add_run(' ').font.size = Pt(1)


def figura(arquivo, legenda, largura_cm=15.5):
    """Imagem centralizada com legenda em itálico."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(IMG + arquivo, width=Cm(largura_cm))
    _legenda(legenda)


def figuras_lado(itens, largura_cm=6.5):
    """Duas ou mais imagens lado a lado (tabela sem bordas); itens = [(arquivo, legenda)]."""
    t = doc.add_table(rows=1, cols=len(itens))
    for cell, (arquivo, legenda) in zip(t.rows[0].cells, itens):
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(IMG + arquivo, width=Cm(largura_cm))
        pl = cell.add_paragraph(style='Caption')
        pl.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _rotulo(pl, 'Figura', legenda)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


# ══════════════════════════════════════════════════════════════════════════════
# CAPA
# ══════════════════════════════════════════════════════════════════════════════
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(80)
r = p.add_run('Plataforma de Bring-up de Hardware')
r.bold = True; r.font.size = Pt(26)
r.font.color.rgb = AZUL_ESCURO

p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run('Guia Completo para Criação de Testes Personalizados')
r2.font.size = Pt(15); r2.bold = True
r2.font.color.rgb = AZUL_MEDIO

doc.add_paragraph()

p3 = doc.add_paragraph()
p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
r3 = p3.add_run('Do zero até o teste rodando na placa — passo a passo')
r3.font.size = Pt(12); r3.italic = True
r3.font.color.rgb = CINZA_TEXTO

doc.add_paragraph()

for line in ['Versao: Sprint 16', 'Alvo: STM32F446RE (NUCLEO-F446RE)',
             'Autor: Rafael Macedo', 'Email: rafael.macedoengenharia@gmail.com']:
    px = doc.add_paragraph()
    px.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rx = px.add_run(line)
    rx.font.size = Pt(10)
    rx.font.color.rgb = CINZA_TEXTO

quebra()


# ══════════════════════════════════════════════════════════════════════════════
# SUMARIO MANUAL
# ══════════════════════════════════════════════════════════════════════════════
titulo_lista('Sumário')
indice('TOC \\o "1-2" \\h \\z \\u',
       'Sumário — abra no Word e pressione F9 (ou rode DOC/_gerar_word.py no Windows).')
quebra()

titulo_lista('Lista de Figuras')
indice('TOC \\h \\z \\c "Figura"', 'Lista de figuras — atualize o campo (F9).')

titulo_lista('Lista de Tabelas')
indice('TOC \\h \\z \\c "Tabela"', 'Lista de tabelas — atualize o campo (F9).')
quebra()


# ══════════════════════════════════════════════════════════════════════════════
# 1. O QUE E A PLATAFORMA
# ══════════════════════════════════════════════════════════════════════════════
h1('1. O que é a Plataforma de Bring-up?')

para(
    'Imagine que você acabou de receber uma placa de circuito impresso recém-fabricada. '
    'Você não sabe se os periféricos estão funcionando corretamente, se os pinos estão '
    'conectados como esperado, ou se o microcontrolador responde ao depurador. '
    'O processo de verificar tudo isso se chama bring-up.',
    space_after=8
)
para(
    'Esta plataforma automatiza esse processo. Ela se conecta ao microcontrolador '
    'via ST-Link (o mesmo chip USB que existe na placa NUCLEO) usando o protocolo SWD '
    '(Serial Wire Debug), e executa testes Python que leem e escrevem diretamente nos '
    'registradores do STM32 — sem precisar compilar nenhum firmware de teste. '
    'A exceção são os testes de alta velocidade (seção 12), que carregam um '
    'pequeno programa na RAM do chip.',
    space_after=8
)

box(
    'Analogia: pense na plataforma como um multímetro programável.\n'
    'Você escreve em Python "verifique se o pino PA5 está configurado como saída"\n'
    'e a plataforma vai lá no chip e te diz se é verdade ou não.',
    fill='E8F4F8',
    title='  Analogia',
    title_color=AZUL_MEDIO
)

h2('O que você consegue fazer com ela:')
bullet('Verificar IDs de dispositivos (confirmar que é o chip certo)')
bullet('Checar configurações de clock e periféricos')
bullet('Testar GPIO: ligar LED, ler botão, checar direção de pinos')
bullet('Fazer loopback de USART e SPI (enviar e receber de volta)')
bullet('Medir tensão interna de referência do ADC')
bullet('Varrer o barramento I2C em busca de dispositivos')
bullet('Escrever e ler de volta a SRAM (teste de integridade de memória)')
bullet('Criar qualquer teste customizado para seu próprio hardware')
bullet('Suportar múltiplos microcontroladores com apenas um arquivo de configuração')
bullet('Identificar automaticamente o microcontrolador conectado ao probe, sem parar o firmware')
bullet('Importar o pinout de um projeto STM32CubeMX (.ioc) ao criar um projeto')
bullet('Medir a velocidade máxima de barramentos com um agente que roda na RAM do MCU')
bullet('Usar osciloscópio e analisador de espectro nos testes, com rastreabilidade de calibração no relatório')
bullet('Montar a suite com blocos reutilizáveis (UART, SPI, I2C, CAN, Ethernet, interrupções, GPIO) escolhendo pinos e periféricos no Plano de testes')
bullet('Avaliar tolerâncias pelo perfil de critérios, com margem e origem de cada limite no relatório')

quebra()


# ══════════════════════════════════════════════════════════════════════════════
# 2. COMO FUNCIONA POR DENTRO
# ══════════════════════════════════════════════════════════════════════════════
h1('2. Como a Plataforma Funciona por Dentro')

para(
    'A plataforma é organizada em camadas. Cada camada tem uma responsabilidade clara. '
    'Você, como criador de testes, trabalha apenas na camada mais alta — o seu teste Python. '
    'As camadas abaixo são gerenciadas automaticamente.',
    space_after=8
)

figura('diagrama_camadas.png',
       'Camadas da plataforma, do seu teste ao microcontrolador')

para(
    'Na coluna principal, cada camada só conversa com a vizinha: o teste chama o PAL pelo nome '
    'do registrador, o PAL consulta o arquivo SVD para achar o endereço e o probe faz o acesso '
    'pelo SWD. Os blocos de comunicação passam pela HAL de periféricos (seção 14), que escolhe '
    'os registradores pela família do chip. Em laranja, os dois caminhos que saem dessa pilha: '
    'os instrumentos de bancada '
    '(seção 13), que medem os pinos com pontas de prova, e o agente na RAM (seção 12), um '
    'programa carregado no próprio chip para testes de alta velocidade. O selo no canto de cada '
    'bloco indica a seção do guia que trata dele.',
    space_after=8
)

h2('O ciclo de vida de um teste')
numbered('Você seleciona a placa no ComboBox "Projeto" na GUI')
numbered('A plataforma chama set_active_board(BOARD) com os dados da placa')
numbered('Você clica "Executar Suite" (ou roda o CLI)')
numbered('A plataforma chama  MeuTeste().run(pal)')
numbered('run() chama  _run(pal)  — o método que você escreve')
numbered('Dentro de _run(), você usa  pal.read_field()  e  pal.modify_field()')
numbered('O PAL consulta o SVD para achar o endereço do campo no STM32')
numbered('O pyOCD envia o comando via USB para o ST-Link')
numbered('O ST-Link acessa o registrador do chip via SWD em microssegundos')
numbered('O valor volta para o Python e você decide: _pass() ou _fail()')
numbered('O resultado aparece colorido na tabela da GUI')

doc.add_paragraph()
box(
    'O arquivo SVD (devices/stm32f446/STM32F446.svd) é o mapa completo de todos\n'
    'os registradores do STM32F446. É um XML padrão da ARM que descreve cada\n'
    'periférico, cada registrador e cada campo de bit com seu nome e posição.\n'
    'A plataforma lê esse arquivo na inicialização e usa os nomes como chaves.',
    fill='E8F4F8',
    title='  O arquivo SVD',
    title_color=AZUL_MEDIO
)

quebra()

h2('A interface gráfica (GUI)')

para(
    'Abra a plataforma com python run.py (ou pelo VS Code, configuração "Bring-up GUI"). '
    'A janela tem um painel à esquerda, com três abas, e a tabela de resultados à direita. '
    'A figura abaixo mostra a janela depois de uma execução completa da suite.',
    space_after=6
)

figura('gui_visao_geral.png',
       'Figura 1 — Janela principal após conectar e executar a suite (dados simulados)')

table(
    ['Nº', 'Elemento', 'Para que serve'],
    [
        ['1', 'Projeto', 'Escolhe o projeto (produto). O projeto define a placa base, o pinout e os testes'],
        ['2', 'Probe + Atualizar', 'Lista os probes ST-Link e FT2232H ligados ao USB'],
        ['3', 'Identificar MCU / Conectar', 'Lê o ID do chip sem parar o firmware (seção 3.5) / abre a sessão SWD para os testes'],
        ['4', 'Abas', 'Projeto, Instrumentos (equipamentos de bancada, seção 13) e Testes (suite, plugins e macros)'],
        ['5', 'Resultados', 'Uma linha por teste: verde = PASS, vermelho = FAIL, com tempo e mensagem'],
        ['6', 'Relatório / Exportar PDF / Exportar', 'Relatório HTML, Word ou PDF / PDF direto / resultados em JSON ou CSV. A janela de salvar abre na última pasta usada'],
        ['7', 'Executar Suite [F5]', 'Roda os testes embutidos e os plugins carregados'],
        ['8', 'Barra de status', 'Projeto ativo, probe conectado e placar da suite (ex.: 7/8 PASS 1 FAIL)'],
    ],
    widths=[1.2, 4.3, 10.5],
    titulo='Elementos da janela principal (Figura 2)',
)

box(
    'As capturas de tela foram geradas pelo DOC/_gerar_capturas.py, que abre a GUI\n'
    'com probe e resultados simulados. Os valores (DEV_ID, mensagens, tempos) são ilustrativos;\n'
    'no hardware real os textos seguem o mesmo formato.\n'
    'O diagrama de camadas (Figura 1) é gerado pelo DOC/_gerar_diagramas.py.',
    fill='E8F4F8',
    title='  Sobre as figuras',
    title_color=AZUL_MEDIO
)

quebra()


# ══════════════════════════════════════════════════════════════════════════════
# 3. BOARDCONFIG — PORTABILIDADE
# ══════════════════════════════════════════════════════════════════════════════
h1('3. BoardConfig — Portabilidade Entre Placas')

para(
    'Um dos maiores desafios em plataformas de bring-up é ter que editar o código dos '
    'testes toda vez que se muda de hardware. A plataforma resolve isso com o conceito '
    'de BoardConfig: um único arquivo Python que descreve toda a informação específica '
    'de uma placa. Os testes não conhecem os pinos — eles perguntam ao BoardConfig.',
    space_after=8
)

h2('3.1 O que o BoardConfig descreve')

table(
    ['Campo', 'Tipo', 'Exemplo (NUCLEO-F446RE)'],
    [
        ['project_code',       'str',   '"2180.0720.00"'],
        ['name',               'str',   '"NUCLEO-F446RE"'],
        ['target',             'str',   '"stm32f446retx"  (pyOCD target)'],
        ['svd_path',           'str',   '"devices/stm32f446/STM32F446.svd"'],
        ['dbgmcu_idcode_addr', 'int',   '0xE0042000'],
        ['expected_dev_id',    'int',   '0x421  (STM32F446)'],
        ['sram_base',          'int',   '0x20000000'],
        ['led_port / led_pin', 'str/int', '"GPIOA", 5  (LED LD2)'],
        ['btn_port / btn_pin', 'str/int', '"GPIOC", 13  (botão azul)'],
        ['btn_active_low',     'bool',  'True  (pull-up, pressionado = 0)'],
        ['required_clocks',    'tuple', '(("AHB1ENR","GPIOAEN","GPIOA"),...)'],
        ['usart',              'UsartConfig', 'n=1, TX=PA9/AF7, RX=PA10/AF7'],
        ['spi',                'SpiConfig',   'n=1, SCK=PA5/AF5, MOSI=PA7/AF5'],
    ],
    widths=[4.5, 3, 8.5],
    titulo='Campos do BoardConfig',
)

h2('3.2 O campo project_code — rastreabilidade')

para(
    'Cada BoardConfig tem um campo project_code (ex.: "2180.0720.00"). Ele aparece em:',
    space_after=4
)
bullet('Título da janela GUI: "Bring-up 2180.0720.00 — NUCLEO-F446RE"')
bullet('Arquivo exportado JSON/CSV: campo "project_code" em cada resultado')
bullet('Log de execução CLI')

para(
    'Isso permite rastrear qual configuração de placa estava ativa quando os '
    'resultados foram gerados — fundamental para documentação de qualidade.',
    space_after=8
)

box(
    'Regra prática: um project_code por revisão de hardware.\n'
    'Se a placa ganha um novo conector, crie 2180.0720.01 com replace().\n'
    'Os testes ficam idênticos — apenas o arquivo de config muda.',
    fill='E8F4F8',
    title='  Rastreabilidade de revisão',
    title_color=AZUL_MEDIO
)

h2('3.3 Como criar o arquivo para uma nova placa')

para(
    'Crie o arquivo  src/boards/minha_placa.py  com o conteúdo abaixo. '
    'A GUI detecta automaticamente qualquer arquivo .py nessa pasta que '
    'exponha uma variável BOARD do tipo BoardConfig.',
    space_after=6
)

code_block([
    [CMT('# src/boards/minha_placa.py'), TXT('')],
    [KW('from '), TXT('core.board.config '), KW('import '), TXT('BoardConfig, PinConfig, UsartConfig, SpiConfig')],
    [''],
    [TXT('BOARD = BoardConfig(')],
    [TXT('    '), CMT('# Identificação')],
    [TXT('    project_code = '), STR('"2180.0720.00"'), TXT(',      '), CMT('# código do projeto')],
    [TXT('    name         = '), STR('"Minha Placa XYZ"'), TXT(',')],
    [''],
    [TXT('    '), CMT('# Conexão pyOCD')],
    [TXT('    target   = '), STR('"stm32l476rgtx"'), TXT(',')],
    [TXT('    svd_path = '), STR('"devices/stm32l476/STM32L476.svd"'), TXT(',')],
    [''],
    [TXT('    '), CMT('# Chip identity')],
    [TXT('    dbgmcu_idcode_addr = '), NUM('0xE0042000'), TXT(',')],
    [TXT('    expected_dev_id    = '), NUM('0x415'), TXT(',       '), CMT('# STM32L476')],
    [''],
    [TXT('    '), CMT('# SRAM')],
    [TXT('    sram_base = '), NUM('0x20000000'), TXT(',')],
    [''],
    [TXT('    '), CMT('# LED (pino de saída de status)')],
    [TXT('    led_port        = '), STR('"GPIOB"'), TXT(',')],
    [TXT('    led_pin         = '), NUM('7'), TXT(',')],
    [TXT('    led_clock_reg   = '), STR('"AHB2ENR"'), TXT(',      '), CMT('# L4 usa AHB2')],
    [TXT('    led_clock_field = '), STR('"GPIOBEN"'), TXT(',')],
    [''],
    [TXT('    '), CMT('# Botão de usuário')],
    [TXT('    btn_port        = '), STR('"GPIOC"'), TXT(',')],
    [TXT('    btn_pin         = '), NUM('13'), TXT(',')],
    [TXT('    btn_clock_reg   = '), STR('"AHB2ENR"'), TXT(',')],
    [TXT('    btn_clock_field = '), STR('"GPIOCEN"'), TXT(',')],
    [TXT('    btn_active_low  = '), KW('True'), TXT(',')],
    [''],
    [TXT('    '), CMT('# Clocks que devem estar habilitados pelo firmware')],
    [TXT('    required_clocks = (')],
    [TXT('        ('), STR('"AHB2ENR"'), TXT(', '), STR('"GPIOBEN"'), TXT(', '), STR('"GPIOB"'), TXT('),')],
    [TXT('        ('), STR('"AHB2ENR"'), TXT(', '), STR('"GPIOCEN"'), TXT(', '), STR('"GPIOC"'), TXT('),')],
    [TXT('    ),')],
    [''],
    [TXT('    '), CMT('# USART para loopback')],
    [TXT('    usart = UsartConfig(')],
    [TXT('        n            = '), NUM('1'), TXT(',')],
    [TXT('        tx           = PinConfig(port='), STR('"GPIOA"'), TXT(', pin='), NUM('9'), TXT(',  af='), NUM('7'), TXT('),')],
    [TXT('        rx           = PinConfig(port='), STR('"GPIOA"'), TXT(', pin='), NUM('10'), TXT(', af='), NUM('7'), TXT('),')],
    [TXT('        clock_reg    = '), STR('"APB2ENR"'), TXT(',')],
    [TXT('        clock_field  = '), STR('"USART1EN"'), TXT(',')],
    [TXT('        loopback_baud = '), NUM('9600'), TXT(',')],
    [TXT('    ),')],
    [''],
    [TXT('    '), CMT('# SPI para loopback')],
    [TXT('    spi = SpiConfig(')],
    [TXT('        n           = '), NUM('1'), TXT(',')],
    [TXT('        sck         = PinConfig(port='), STR('"GPIOA"'), TXT(', pin='), NUM('5'), TXT(', af='), NUM('5'), TXT('),')],
    [TXT('        miso        = PinConfig(port='), STR('"GPIOA"'), TXT(', pin='), NUM('6'), TXT(', af='), NUM('5'), TXT('),')],
    [TXT('        mosi        = PinConfig(port='), STR('"GPIOA"'), TXT(', pin='), NUM('7'), TXT(', af='), NUM('5'), TXT('),')],
    [TXT('        clock_reg   = '), STR('"APB2ENR"'), TXT(',')],
    [TXT('        clock_field = '), STR('"SPI1EN"'), TXT(',')],
    [TXT('    ),')],
    [TXT(')')],
], caption='src/boards/minha_placa.py — template completo')

h2('3.4 Usar get_active_board() nos seus testes')

para(
    'Dentro do método _run(), chame get_active_board() para obter os dados '
    'da placa que foi selecionada na GUI. O teste não precisa saber qual placa '
    'está conectada — ele simplesmente pergunta.',
    space_after=6
)

code_block([
    [KW('from '), TXT('tests_builtin.base '), KW('import '), TXT('BuiltinTest, TestResult')],
    [''],
    [KW('class '), DEF('MeuTestePortavel'), TXT('(BuiltinTest):')],
    [TXT('    name = '), STR('"Meu Teste"')],
    [''],
    [KW('    def '), DEF('_run'), TXT('(self, pal) -> TestResult:')],
    [TXT('        '), CMT('# Sempre dentro de _run(), nunca no __init__')],
    [KW('        from '), TXT('core.board.config '), KW('import '), TXT('get_active_board')],
    [TXT('        cfg = get_active_board()')],
    [''],
    [TXT('        '), CMT('# Agora use os campos do BoardConfig')],
    [TXT('        pal.modify_field('), STR('"RCC"'), TXT(', cfg.led_clock_reg, cfg.led_clock_field, '), NUM('1'), TXT(')')],
    [TXT('        nivel = pal.read_field(cfg.led_port, '), STR('"ODR"'), TXT(', f'), STR('"ODR{cfg.led_pin}"'), TXT(')')],
    [''],
    [KW('        if '), TXT('nivel == '), NUM('1'), TXT(':')],
    [KW('            return '), TXT('self._pass('), STR('"LED ON"'), TXT(')')],
    [KW('        return '), TXT('self._fail('), STR('"LED não subiu"'), TXT(')')],
], caption='Usando get_active_board() para portabilidade')

box(
    'Para trocar de placa: basta selecionar outro item no ComboBox "Projeto" na GUI.\n'
    'Nenhuma linha de código de teste precisa mudar.\n'
    'Para adicionar suporte à nova placa: crie src/boards/nova_placa.py com BOARD=BoardConfig(...).',
    fill='E8F4F8',
    title='  Regra de portabilidade',
    title_color=AZUL_MEDIO
)

h2('3.5 Identificar o MCU conectado')

para(
    'Antes de rodar qualquer teste, vale confirmar que o chip na bancada é o mesmo da placa '
    'selecionada no projeto. O botão "Identificar MCU", na aba Projeto, faz isso sem depender '
    'da placa escolhida e sem parar o firmware: ele só lê registradores de identificação.'
)

h3('Como usar')
bullet('Conecte o probe (ST-Link ou FT2232H) e clique "Atualizar" para listá-lo no ComboBox')
bullet('Selecione o probe e clique "Identificar MCU" — não é preciso conectar antes')
bullet('Leia o resultado abaixo dos botões: família, núcleo, DEV_ID/REV_ID e DPIDR')
bullet('Se o chip não for o da placa do projeto, aparece um aviso em laranja com o DEV_ID esperado')

figuras_lado([
    ('gui_mcu_confere.png', 'Figura 2 — Chip confere com a placa do projeto'),
    ('gui_mcu_nao_confere.png', 'Figura 3 — Chip diferente (STM32H7): aviso em laranja'),
])

code(
    'MCU: STM32H74x/H75x\n'
    'DEV_ID=0x450  REV_ID=0x2003\n'
    'DPIDR=0x6BA02477\n'
    '! Não confere com a placa do projeto (NUCLEO-F446RE):\n'
    '  esperado DEV_ID 0x421 (STM32F446)',
    caption='Exemplo real: STLINK-V3 ligado a uma placa com STM32H7, projeto da NUCLEO-F446RE'
)

h3('O que é lido')
table(
    ['Registrador', 'Endereço', 'O que informa'],
    [
        ['DPIDR', 'DP registrador 0', 'Versão do Debug Port (sempre legível)'],
        ['CPUID', '0xE000ED00', 'Núcleo ARM: Cortex-M0, M3, M4, M7, M33...'],
        ['ROM table PIDR', '0xE00FFFD0..FEC', 'Fabricante do chip (código JEP106)'],
        ['DBGMCU_IDCODE', '0xE0042000 (F1/F2/F3/F4/F7/L1/L4/G4)', 'DEV_ID (família) e REV_ID (revisão)'],
        ['DBGMCU_IDCODE', '0x40015800 (F0/L0/G0/C0)', 'Idem, famílias com Cortex-M0/M0+'],
        ['DBGMCU_IDCODE', '0xE00E1000 via AP2 (H7)', 'Idem; no H7 o DBGMCU só é visível pela AP2'],
    ],
    widths=[4.5, 4.5, 7],
    titulo='Registradores lidos na identificação do MCU',
)

para(
    'O programa lê as portas de acesso (APs) do chip diretamente, sem abrir sessão de depuração '
    'com o núcleo. Por isso a identificação funciona mesmo quando o núcleo não responde (modo de '
    'baixo consumo, travado), situação em que o pyOCD retorna "No cores were discovered".'
)

code_block([
    [KW('from '), TXT('core.target '), KW('import '), TXT('identify')],
    '',
    [KW('def '), DEF('_run'), TXT('(self, pal) -> TestResult:')],
    [TXT('    info = identify(pal.mem, dpidr='), NUM('0'), TXT(')   '), CMT('# dpidr só aparece no resumo')],
    [KW('    if '), TXT('info.family '), KW('is None'), TXT(':')],
    [KW('        return '), TXT('self._skip('), STR('"Chip fora da tabela STM32"'), TXT(')')],
    [KW('    return '), TXT('self._pass(info.family, measured=info.dev_id)')],
], caption='Usando a identificação dentro de um teste')

box(
    'A tabela DEV_ID -> família fica em src/core/target/identify.py e deve ser conferida\n'
    'com o manual de referência (RM) do chip antes de servir de critério de aprovação.\n'
    'FT2232H: identificação implementada mas ainda não testada com hardware; lê apenas\n'
    'pela AP0 (não identifica o STM32H7). O caminho FT2232H usa SWD, não a cadeia JTAG.',
    fill='E8F4F8',
    title='  Limitações',
    title_color=AZUL_MEDIO
)

h2('3.6 Importar o pinout do STM32CubeMX (.ioc)')

para(
    'Se o firmware da placa foi configurado no STM32CubeMX, o arquivo .ioc já contém o pinout. '
    'O editor de projeto importa esse arquivo e preenche a tabela de sinais, evitando digitar '
    'pino por pino.'
)

h3('Como usar')
bullet('Na aba Projeto, clique "Novo..." (ou "Editar" em um projeto existente)')
bullet('Na aba Pinout do editor, clique "Importar .ioc..." e escolha o arquivo')
bullet('Confira a janela de confirmação: MCU, número de pinos, placa base e pinos ignorados')
bullet('Confirme — a tabela atual é substituída pelos pinos do .ioc')
bullet('Revise nomes, descrições e a coluna "Ativo baixo" e clique "Salvar"')

figura('gui_ioc_confirmacao.png',
       'Figura 4 — Janela de confirmação: MCU, quantidade de pinos, placa base e pinos ignorados',
       largura_cm=11)

figura('gui_ioc_pinout.png',
       'Figura 5 — Editor de projeto após a importação: (1) botão "Importar .ioc..."; '
       '(2) tabela de sinais preenchida')

h3('Como cada pino é convertido')
table(
    ['No .ioc', 'No projeto', 'Observação'],
    [
        ['PC4.GPIO_Label=LED1 / GPIO_Output', 'led1 — GPIO', 'O rótulo do CubeMX vira o nome do sinal'],
        ['PD5.Signal=USART2_TX', 'usart2_tx — UART USART2 TX', 'Baud rate incluído quando gravado no .ioc'],
        ['PB5.Signal=GPXTI5 (IT_FALLING)', 'pb5_exti — EXTI EXTI5 FALLING', 'Pull-up/pull-down vai para a descrição'],
        ['PF10.Signal=SharedAnalog_PF10', 'adc3_inp6 — ADC ADC3 INP6', 'Resolvido pelas linhas SH.*; prefere entrada INP'],
        ['TIM1_CH4 com PWM Generation', 'tim1_ch4 — PWM', 'Sem PWM configurado, fica como TIM'],
        ['USB_OTG_FS_DP, ETH_..., SDMMC_...', 'GPIO', 'Sem tipo equivalente; sinal original na descrição'],
        ['SYS_JTMS-SWDIO, RCC_OSC_IN...', '(ignorado)', 'Pinos de debug e oscilador não são sinais do produto'],
    ],
    widths=[4.5, 4.5, 7],
    titulo='Conversão dos pinos do .ioc para o projeto',
)

para(
    'A placa base é escolhida automaticamente quando existe um BoardConfig cujo campo target '
    'corresponde ao MCU do .ioc (Mcu.UserName em minúsculas, ex.: STM32F446RETx -> stm32f446retx). '
    'Se não existir, o projeto fica com a placa selecionada no formulário e a janela avisa.'
)

code_block([
    [KW('from '), TXT('core.project.config '), KW('import '), TXT('get_active_project')],
    '',
    [TXT('tx = get_active_project().pin('), STR('"usart2_tx"'), TXT(')   '), CMT('# GPIOD.5, USART2 TX')],
    [TXT('uart = UsartInstrument(tx.peripheral.name, pal)')],
], caption='Usando um sinal importado dentro de um teste')

box(
    'O .ioc não informa se um sinal é ativo em nível baixo: a coluna "Ativo baixo"\n'
    'vem desmarcada e deve ser conferida com o esquemático.\n'
    'Para um MCU sem placa cadastrada (ex.: STM32H753), crie src/boards/<placa>.py com\n'
    'target igual ao Mcu.UserName em minúsculas (ex.: target="stm32h753zitx").',
    fill='E8F4F8',
    title='  Revise antes de salvar',
    title_color=AZUL_MEDIO
)

h2('3.7 Placas STM32H7 (NUCLEO-H753ZI)')

para(
    'A NUCLEO-H753ZI (src/boards/nucleo_h753zi.py) mostra o que muda numa família diferente do '
    'F4. Os testes não mudam: as diferenças ficam no BoardConfig e no SVD.',
    space_after=6
)
table(
    ['Campo do BoardConfig', 'NUCLEO-H753ZI', 'Para que serve'],
    [
        ['target', 'stm32h753zitx', 'Igual ao Mcu.UserName do CubeMX (importação do .ioc)'],
        ['pyocd_target', 'stm32h743xx', 'Alvo embutido no pyOCD quando difere de target'],
        ['connect_mode', 'under-reset', 'Conecta com o reset pressionado (firmware em baixo consumo)'],
        ['resume_after_connect', 'True', 'Libera o firmware logo após conectar'],
        ['gpio_clock_reg', 'AHB4ENR', 'Registrador dos bits GPIOxEN (AHB1ENR no F4)'],
        ['usart / spi', 'None', 'Loopbacks embutidos do F4 viram SKIP; use os blocos (seção 14)'],
        ['dbgmcu_idcode_addr', '0x5C001000', 'DBGMCU_IDCODE do H7'],
    ],
    widths=[4.0, 3.5, 8.5],
    titulo='Campos do BoardConfig específicos da família H7',
)
para(
    'O SVD do H7 usa nomes novos para os campos de GPIO (MODE5, OD5, ID5, AFSEL8). O PAL aceita '
    'também os nomes do F4 (MODER5, ODR5, IDR5, AFRH8), então um teste escrito para uma família '
    'roda na outra sem alteração.',
    space_after=6
)
box(
    'Conexão "under-reset": o pyOCD liga os bits de debug em baixo consumo (DBGMCU_CR) e o\n'
    'firmware volta a rodar. A descoberta do chip às vezes falha logo após uma sessão anterior;\n'
    'a plataforma tenta até 3 vezes antes de acusar erro.',
    fill='E8F4F8',
    title='  Conexão no H7',
    title_color=AZUL_MEDIO
)

quebra()


# ══════════════════════════════════════════════════════════════════════════════
# 4. ANATOMIA DE UM TESTE
# ══════════════════════════════════════════════════════════════════════════════
h1('4. Anatomia de um Teste — Padrões A e B')

para(
    'Existem dois padrões de teste na plataforma. Entender qual usar em cada '
    'situação é fundamental para escrever testes corretos e portáveis.',
    space_after=10
)

h2('Padrão A — Teste Portável (usa BoardConfig)')

para(
    'Use quando o teste exercita hardware da placa: LED, botão, USART para loopback, '
    'pinos de GPIO. Ao trocar de placa na GUI, o teste se adapta automaticamente '
    'sem nenhuma alteração de código.',
    space_after=6
)

code_block([
    [KW('from '), TXT('tests_builtin.base '), KW('import '), TXT('BuiltinTest, TestResult')],
    [''],
    [KW('class '), DEF('TestePortavel'), TXT('(BuiltinTest):')],
    [TXT('    name        = '), STR('"Teste do LED"')],
    [TXT('    description = '), STR('"Verifica o LED da placa ativa — portável entre placas"')],
    [''],
    [KW('    def '), DEF('_run'), TXT('(self, pal) -> TestResult:')],
    [TXT('        '), CMT('# 1. Obtém a configuração da placa selecionada na GUI')],
    [KW('        from '), TXT('core.board.config '), KW('import '), TXT('get_active_board')],
    [TXT('        cfg = get_active_board()')],
    [''],
    [TXT('        '), CMT('# 2. Usa os campos — funciona para qualquer placa')],
    [TXT('        pal.modify_field('), STR('"RCC"'), TXT(', cfg.led_clock_reg, cfg.led_clock_field, '), NUM('1'), TXT(')')],
    [TXT('        orig = pal.read_field(cfg.led_port, '), STR('"MODER"'), TXT(', f'), STR('"MODER{cfg.led_pin}"'), TXT(')')],
    [''],
    [KW('        try'), TXT(':')],
    [TXT('            pal.modify_field(cfg.led_port, '), STR('"MODER"'), TXT(', f'), STR('"MODER{cfg.led_pin}"'), TXT(', '), NUM('0b01'), TXT(')')],
    [TXT('            pal.write_reg(cfg.led_port, '), STR('"BSRR"'), TXT(', '), NUM('1'), TXT(' << cfg.led_pin)')],
    [TXT('            nivel = pal.read_field(cfg.led_port, '), STR('"ODR"'), TXT(', f'), STR('"ODR{cfg.led_pin}"'), TXT(')')],
    [''],
    [KW('            if '), TXT('nivel == '), NUM('1'), TXT(':')],
    [KW('                return '), TXT('self._pass('), STR('"LED ON: OK"'), TXT(', measured=nivel, expected='), NUM('1'), TXT(')')],
    [KW('            return '), TXT('self._fail('), STR('"LED não subiu"'), TXT(', measured=nivel, expected='), NUM('1'), TXT(')')],
    [''],
    [KW('        finally'), TXT(':')],
    [TXT('            pal.modify_field(cfg.led_port, '), STR('"MODER"'), TXT(', f'), STR('"MODER{cfg.led_pin}"'), TXT(', orig)')],
], caption='Padrão A: usa get_active_board() — portável entre placas')

h2('Padrão B — Constantes Fixas (verifica firmware)')

para(
    'Use quando o teste verifica parâmetros configurados pelo firmware: baud rate, '
    'frequência de PWM, canal DMA, modo de SPI. Esses valores são da aplicação, '
    'não da placa física — portanto ficam como constantes no topo do arquivo.',
    space_after=6
)

code_block([
    [KW('from '), TXT('tests_builtin.base '), KW('import '), TXT('BuiltinTest, TestResult')],
    [''],
    [CMT('# ─── Constantes da aplicação — altere conforme o firmware ──────────────────')],
    [TXT('PERIFERICO  = '), STR('"USART2"')],
    [TXT('REGISTRADOR = '), STR('"BRR"')],
    [TXT('CAMPO       = '), STR('"DIV_Mantissa"')],
    [TXT('VALOR_OK    = '), NUM('8'), TXT('          '), CMT('# BRR mant para 115200 @ 16 MHz')],
    [CMT('# ────────────────────────────────────────────────────────────────────────────')],
    [''],
    [KW('class '), DEF('UsartBaudCheck'), TXT('(BuiltinTest):')],
    [TXT('    name        = '), STR('"USART2 Baud Check"')],
    [TXT('    description = '), STR('"Verifica se firmware configurou USART2 para 115200 baud"')],
    [''],
    [KW('    def '), DEF('_run'), TXT('(self, pal) -> TestResult:')],
    [TXT('        valor = pal.read_field(PERIFERICO, REGISTRADOR, CAMPO)')],
    [''],
    [KW('        if '), TXT('valor == VALOR_OK:')],
    [KW('            return '), TXT('self._pass(')],
    [TXT('                '), STR('f"{PERIFERICO}.{CAMPO} = {valor} (OK)"'), TXT(',')],
    [TXT('                measured=valor, expected=VALOR_OK')],
    [TXT('            )')],
    [KW('        return '), TXT('self._fail(')],
    [TXT('            '), STR('f"{PERIFERICO}.{CAMPO} = {valor}, esperado {VALOR_OK}"'), TXT(',')],
    [TXT('            measured=valor, expected=VALOR_OK')],
    [TXT('        )')],
], caption='Padrão B: constantes fixas — verifica configuração do firmware')

table(
    ['Critério', 'Padrão A (portável)', 'Padrão B (fixo)'],
    [
        ['O que verifica',   'Hardware da placa (LED, botão, loopback)', 'Firmware da aplicação (baud, PWM, DMA)'],
        ['Como obtém pinos', 'get_active_board() dentro de _run()',       'Constantes no topo do arquivo'],
        ['Ao trocar de placa', 'Funciona automaticamente',               'Constantes continuam as mesmas'],
        ['Quando usar',      'Teste de trazer-a-vida (bring-up)',        'Teste de configuração de firmware'],
    ],
    widths=[4, 6, 6],
    titulo='Comparação entre os Padrões A e B',
)

quebra()


# ══════════════════════════════════════════════════════════════════════════════
# 5. O PAL
# ══════════════════════════════════════════════════════════════════════════════
h1('5. O PAL: sua interface com o hardware')

para(
    'O PAL (Peripheral Access Layer) é o objeto que você recebe como parâmetro "pal" '
    'no método _run(). Pense nele como um dicionário gigante que mapeia nomes em '
    'ações de hardware. Você não precisa saber qual é o endereço de memória do '
    'registrador MODER do GPIOA — o PAL descobre isso sozinho usando o SVD.',
    space_after=10
)

h2('5.1 read_field() — Ler um campo de registrador')

para(
    'Esta é a operação mais comum. Você informa o nome do periférico, do registrador '
    'e do campo, e recebe o valor numérico daquele campo específico.',
    space_after=6
)

code_block([
    [CMT('# Sintaxe:'), TXT('')],
    [TXT('valor = pal.read_field('), STR('"PERIFERICO"'), TXT(', '), STR('"REGISTRADOR"'), TXT(', '), STR('"CAMPO"'), TXT(')')],
    [''],
    [CMT('# Exemplos práticos:'), TXT('')],
    [''],
    [CMT('# Verifica se o clock do GPIOA está habilitado (campo de 1 bit):')],
    [TXT('clock_on = pal.read_field('), STR('"RCC"'), TXT(', '), STR('"AHB1ENR"'), TXT(', '), STR('"GPIOAEN"'), TXT(')')],
    [CMT('# clock_on será 0 ou 1')],
    [''],
    [CMT('# Lê o modo do pino 5 do GPIOA (campo de 2 bits):')],
    [TXT('modo = pal.read_field('), STR('"GPIOA"'), TXT(', '), STR('"MODER"'), TXT(', '), STR('"MODER5"'), TXT(')')],
    [CMT('# modo será: 0=entrada, 1=saída, 2=AF, 3=analógico')],
    [''],
    [CMT('# Lê o estado do botão (pino PC13):')],
    [TXT('botao = pal.read_field('), STR('"GPIOC"'), TXT(', '), STR('"IDR"'), TXT(', '), STR('"IDR13"'), TXT(')')],
    [CMT('# botão será 1 (solto) ou 0 (pressionado, com pull-up)')],
])

box(
    'O PAL já faz a máscara e o deslocamento para você. Se o campo MODER5 fica\n'
    'nos bits [11:10] do registrador MODER, o PAL extrai esses 2 bits e retorna\n'
    'um valor entre 0 e 3 — você não precisa fazer "(valor >> 10) & 0x3" na mão.',
    fill='E8F4F8',
    title='  Sem contas de bit',
    title_color=AZUL_MEDIO
)

h2('5.2 modify_field() — Modificar um campo (Read-Modify-Write)')

para(
    'Esta operação lê o registrador inteiro, modifica apenas os bits do campo '
    'especificado e escreve o valor de volta. Os outros campos do registrador '
    'não são alterados.',
    space_after=6
)

code_block([
    [CMT('# Sintaxe:'), TXT('')],
    [TXT('pal.modify_field('), STR('"PERIFERICO"'), TXT(', '), STR('"REGISTRADOR"'), TXT(', '), STR('"CAMPO"'), TXT(', novo_valor)')],
    [''],
    [CMT('# Exemplos:'), TXT('')],
    [''],
    [CMT('# Configura PA5 como saída (MODER5 = 0b01):')],
    [TXT('pal.modify_field('), STR('"GPIOA"'), TXT(', '), STR('"MODER"'), TXT(', '), STR('"MODER5"'), TXT(', '), NUM('0b01'), TXT(')')],
    [''],
    [CMT('# Habilita o clock do GPIOA:')],
    [TXT('pal.modify_field('), STR('"RCC"'), TXT(', '), STR('"AHB1ENR"'), TXT(', '), STR('"GPIOAEN"'), TXT(', '), NUM('1'), TXT(')')],
    [''],
    [CMT('# Ou usando campos do BoardConfig (portável):')],
    [TXT('cfg = get_active_board()')],
    [TXT('pal.modify_field('), STR('"RCC"'), TXT(', cfg.led_clock_reg, cfg.led_clock_field, '), NUM('1'), TXT(')')],
])

h2('5.3 read_reg() e write_reg() — Registrador Inteiro')

para(
    'Quando você precisa ler ou escrever o registrador completo de 32 bits de uma vez. '
    'Útil para registradores de status que são lidos de uma vez (como USART SR) '
    'ou para o BSRR do GPIO (Bit Set/Reset Register).',
    space_after=6
)

code_block([
    [CMT('# Ler o registrador de status da USART1 por inteiro:')],
    [TXT('sr = pal.read_reg('), STR('"USART1"'), TXT(', '), STR('"SR"'), TXT(')')],
    [TXT('txe  = (sr >> '), NUM('7'), TXT(') & '), NUM('1'), TXT('   '), CMT('# extrai bit 7 manualmente')],
    [''],
    [CMT('# Escrever no BSRR para controlar pinos de saída:')],
    [TXT('pal.write_reg('), STR('"GPIOA"'), TXT(', '), STR('"BSRR"'), TXT(', '), NUM('1'), TXT(' << '), NUM('5'), TXT(')    '), CMT('# SET PA5')],
    [TXT('pal.write_reg('), STR('"GPIOA"'), TXT(', '), STR('"BSRR"'), TXT(', '), NUM('1'), TXT(' << ('), NUM('5'), TXT(' + '), NUM('16'), TXT('))'), CMT('# RESET PA5')],
])

h2('5.4 pal.mem — Acesso Direto por Endereço')

para(
    'Para situações especiais onde você precisa acessar um endereço de memória que '
    'não está no SVD (como registradores de debug da ARM ou regiões de SRAM).',
    space_after=6
)

code_block([
    [CMT('# Ler o ID do dispositivo no registrador DBGMCU_IDCODE:')],
    [TXT('idcode = pal.mem.read_mem32('), NUM('0xE0042000'), TXT(')')],
    [TXT('dev_id = idcode & '), NUM('0xFFF'), TXT('   '), CMT('# bits [11:0] = device ID')],
    [''],
    [CMT('# Ou usando o BoardConfig (portável):')],
    [TXT('cfg    = get_active_board()')],
    [TXT('idcode = pal.mem.read_mem32(cfg.dbgmcu_idcode_addr)')],
    [TXT('dev_id = idcode & '), NUM('0xFFF')],
    [''],
    [CMT('# Escrever na SRAM (para teste de memória):')],
    [TXT('pal.mem.write_mem32(cfg.sram_base, '), NUM('0xA5A5A5A5'), TXT(')')],
    [TXT('lido = pal.mem.read_mem32(cfg.sram_base)')],
])

quebra()


# ══════════════════════════════════════════════════════════════════════════════
# 6. DESCOBRINDO NOMES
# ══════════════════════════════════════════════════════════════════════════════
h1('6. Como Descobrir Nomes de Periféricos e Campos')

para(
    'A pergunta mais comum de quem começa é: "Como eu sei o nome exato do periférico, '
    'do registrador e do campo que eu preciso usar?" A resposta é: consultando o SVD.',
    space_after=8
)

h2('6.1 Consultando via terminal (mais rápido)')

para('Abra o terminal do VS Code (Ctrl+`) e use os comandos abaixo:')

code_block([
    [CMT('# Primeiro, configure o PYTHONPATH:')],
    [TXT('$env:PYTHONPATH = '), STR('"src"')],
    [''],
    [CMT('# Ver todos os periféricos disponíveis:')],
    [TXT('.venv\\Scripts\\python.exe -c "'), TXT('')],
    [TXT('from core.svd.parser import parse; from pathlib import Path')],
    [TXT('svd = parse(Path(\'devices/stm32f446/STM32F446.svd\'))')],
    [TXT('print([p.name for p in svd.peripherals])')],
    [TXT('"')],
    [''],
    [CMT('# Listar registradores de um periférico (ex: USART1):')],
    [TXT('.venv\\Scripts\\python.exe -c "'), TXT('')],
    [TXT('from core.svd.parser import parse; from pathlib import Path')],
    [TXT('svd = parse(Path(\'devices/stm32f446/STM32F446.svd\'))')],
    [TXT('for r in svd.peripheral(\'USART1\').registers:')],
    [TXT('    print(f\'  {r.name:12} offset=0x{r.offset:02X}\')')],
    [TXT('"')],
    [''],
    [CMT('# Listar campos de um registrador (ex: USART1.CR1):')],
    [TXT('.venv\\Scripts\\python.exe -c "'), TXT('')],
    [TXT('from core.svd.parser import parse; from pathlib import Path')],
    [TXT('svd = parse(Path(\'devices/stm32f446/STM32F446.svd\'))')],
    [TXT('for f in svd.peripheral(\'USART1\').register(\'CR1\').fields:')],
    [TXT('    print(f\'  {f.name:12} bit[{f.bit_offset}:{f.bit_offset+f.bit_width-1}]\')')],
    [TXT('"')],
])

h2('6.2 Periféricos mais usados e seus clocks')

table(
    ['Periférico', 'Clock em', 'Campo no RCC'],
    [
        ['GPIOA, GPIOB, GPIOC...', 'AHB1ENR', 'GPIOAEN, GPIOBEN, GPIOCEN...'],
        ['USART1, USART6', 'APB2ENR', 'USART1EN, USART6EN'],
        ['USART2, USART3, UART4, UART5', 'APB1ENR', 'USART2EN, USART3EN, UART4EN, UART5EN'],
        ['SPI1, SPI4', 'APB2ENR', 'SPI1EN, SPI4EN'],
        ['SPI2, SPI3', 'APB1ENR', 'SPI2EN, SPI3EN'],
        ['I2C1, I2C2, I2C3', 'APB1ENR', 'I2C1EN, I2C2EN, I2C3EN'],
        ['TIM1, TIM8', 'APB2ENR', 'TIM1EN, TIM8EN'],
        ['TIM2..TIM7', 'APB1ENR', 'TIM2EN, TIM3EN...'],
        ['ADC1, ADC2, ADC3', 'APB2ENR', 'ADC1EN, ADC2EN, ADC3EN'],
        ['DMA1, DMA2', 'AHB1ENR', 'DMA1EN, DMA2EN'],
    ],
    widths=[5, 3.5, 7.5],
    titulo='Periféricos mais usados e seus clocks (STM32F446)',
)

h2('6.3 Alternate Functions mais usadas (NUCLEO-F446RE)')

table(
    ['Sinal', 'Pino', 'Label CN9', 'Campo AFR', 'Valor AF'],
    [
        ['USART1_TX', 'PA9',  'D8',  'GPIOA.AFRH.AFRH9',   'AF7'],
        ['USART1_RX', 'PA10', 'D2',  'GPIOA.AFRH.AFRH10',  'AF7'],
        ['SPI1_SCK',  'PA5',  'D13', 'GPIOA.AFRL.AFRL5',   'AF5'],
        ['SPI1_MISO', 'PA6',  'D12', 'GPIOA.AFRL.AFRL6',   'AF5'],
        ['SPI1_MOSI', 'PA7',  'D11', 'GPIOA.AFRL.AFRL7',   'AF5'],
        ['I2C1_SCL',  'PB8',  'D15', 'GPIOB.AFRH.AFRH8',   'AF4'],
        ['I2C1_SDA',  'PB9',  'D14', 'GPIOB.AFRH.AFRH9',   'AF4'],
        ['TIM3_CH1',  'PA6',  'D12', 'GPIOA.AFRL.AFRL6',   'AF2'],
    ],
    widths=[3, 2, 2.5, 5, 2.5],
    titulo='Alternate Functions mais usadas (NUCLEO-F446RE)',
)

quebra()


# ══════════════════════════════════════════════════════════════════════════════
# 7. PADROES ESSENCIAIS
# ══════════════════════════════════════════════════════════════════════════════
h1('7. Padrões Essenciais de Implementação')

h2('7.1 Sempre Habilite o Clock Antes de Acessar um Periférico')

para(
    'No STM32, todo periférico tem seu clock desabilitado por padrão para economizar '
    'energia. Se você tentar ler ou escrever em um registrador de periférico sem '
    'habilitar o clock, o valor lido será sempre 0 e escritas serão ignoradas '
    '— o que leva a testes que passam mas não testam nada de verdade.',
    space_after=6
)

code_block([
    [CMT('# ERRADO — não habilita clock antes de usar GPIOA:')],
    [KW('def '), DEF('_run'), TXT('(self, pal):')],
    [TXT('    modo = pal.read_field('), STR('"GPIOA"'), TXT(', '), STR('"MODER"'), TXT(', '), STR('"MODER5"'), TXT(')')],
    [TXT('    '), CMT('# modo sempre será 0 se o clock não estiver ligado!')],
    [''],
    [CMT('# CORRETO — habilita clock primeiro:')],
    [KW('def '), DEF('_run'), TXT('(self, pal):')],
    [TXT('    pal.modify_field('), STR('"RCC"'), TXT(', '), STR('"AHB1ENR"'), TXT(', '), STR('"GPIOAEN"'), TXT(', '), NUM('1'), TXT(')')],
    [TXT('    modo = pal.read_field('), STR('"GPIOA"'), TXT(', '), STR('"MODER"'), TXT(', '), STR('"MODER5"'), TXT(')')],
    [TXT('    '), CMT('# agora sim o valor é real')],
])

h2('7.2 O Padrão Salvar / Modificar / Restaurar')

para(
    'Este é o padrão mais importante de toda a plataforma. Se o seu teste '
    'modifica registradores, ele DEVE restaurar o estado original ao terminar '
    '— independentemente de ter passado ou falhado. Isso garante que os próximos '
    'testes não sejam afetados. Use sempre o bloco try/finally.',
    space_after=6
)

code_block([
    [KW('def '), DEF('_run'), TXT('(self, pal):')],
    [''],
    [TXT('    '), CMT('# PASSO 1: salve o estado original ANTES de qualquer alteração')],
    [TXT('    estado_moder5 = pal.read_field('), STR('"GPIOA"'), TXT(', '), STR('"MODER"'), TXT(', '), STR('"MODER5"'), TXT(')')],
    [TXT('    estado_odr5   = pal.read_field('), STR('"GPIOA"'), TXT(', '), STR('"ODR"'), TXT(',   '), STR('"ODR5"'), TXT(')')],
    [''],
    [KW('    try'), TXT(':')],
    [TXT('        '), CMT('# PASSO 2: faça suas alterações e testes aqui')],
    [TXT('        pal.modify_field('), STR('"GPIOA"'), TXT(', '), STR('"MODER"'), TXT(', '), STR('"MODER5"'), TXT(', '), NUM('0b01'), TXT(')')],
    [TXT('        pal.write_reg('), STR('"GPIOA"'), TXT(', '), STR('"BSRR"'), TXT(', '), NUM('1'), TXT(' << '), NUM('5'), TXT(')')],
    [TXT('        nivel = pal.read_field('), STR('"GPIOA"'), TXT(', '), STR('"ODR"'), TXT(', '), STR('"ODR5"'), TXT(')')],
    [''],
    [KW('        if '), TXT('nivel == '), NUM('1'), TXT(':')],
    [KW('            return '), TXT('self._pass('), STR('"PA5 foi para HIGH"'), TXT(')')],
    [KW('        return '), TXT('self._fail('), STR('"PA5 não subiu para HIGH"'), TXT(', measured=nivel, expected='), NUM('1'), TXT(')')],
    [''],
    [KW('    finally'), TXT(':')],
    [TXT('        '), CMT('# PASSO 3: restaure SEMPRE — mesmo se der exceção acima')],
    [TXT('        pal.modify_field('), STR('"GPIOA"'), TXT(', '), STR('"MODER"'), TXT(', '), STR('"MODER5"'), TXT(', estado_moder5)')],
    [TXT('        pal.modify_field('), STR('"GPIOA"'), TXT(', '), STR('"ODR"'), TXT(',   '), STR('"ODR5"'), TXT(',   estado_odr5)')],
])

h2('7.3 Aguardar uma Condição com Timeout')

para(
    'Periféricos como USART e SPI não ficam prontos instantaneamente. Você precisa '
    'esperar que certas flags se tornem 1 antes de prosseguir. Sempre coloque um '
    'timeout para evitar loop infinito caso algo de errado.',
    space_after=6
)

code_block([
    [KW('import '), TXT('time')],
    [''],
    [KW('def '), DEF('_aguardar_flag'), TXT('(pal, periferico, reg, campo, esperado='), NUM('1'), TXT(', timeout_ms='), NUM('100'), TXT('):')],
    [TXT('    prazo = time.perf_counter() + timeout_ms / '), NUM('1000')],
    [KW('    while '), TXT('pal.read_field(periferico, reg, campo) != esperado:')],
    [KW('        if '), TXT('time.perf_counter() > prazo:')],
    [KW('            return '), KW('False')],
    [KW('    return '), KW('True')],
    [''],
    [CMT('# Usando no seu teste:')],
    [KW('def '), DEF('_run'), TXT('(self, pal):')],
    [TXT('    pal.write_reg('), STR('"USART1"'), TXT(', '), STR('"DR"'), TXT(', '), NUM('0xAA'), TXT(')')],
    [KW('    if not '), TXT('_aguardar_flag(pal, '), STR('"USART1"'), TXT(', '), STR('"SR"'), TXT(', '), STR('"RXNE"'), TXT(', timeout_ms='), NUM('200'), TXT('):')],
    [KW('        return '), TXT('self._fail('), STR('"Timeout! Verifique o fio entre TX e RX"'), TXT(')')],
    [TXT('    eco = pal.read_reg('), STR('"USART1"'), TXT(', '), STR('"DR"'), TXT(') & '), NUM('0xFF')],
])

quebra()


# ══════════════════════════════════════════════════════════════════════════════
# 8. EXEMPLOS COMPLETOS
# ══════════════════════════════════════════════════════════════════════════════
h1('8. Exemplos Completos por Periférico')

h2('8.1 Exemplo: Teste de LED Portável (Padrão A)')

para(
    'Objetivo: configurar o pino do LED da placa ativa como saída, ligar, verificar '
    'que foi para nível alto, apagar e verificar nível baixo. Funciona em qualquer '
    'placa configurada no BoardConfig.',
    space_after=6
)

code_block([
    [KW('from '), TXT('tests_builtin.base '), KW('import '), TXT('BuiltinTest, TestResult')],
    [''],
    [KW('class '), DEF('LedVerificacaoTest'), TXT('(BuiltinTest):')],
    [TXT('    name        = '), STR('"LED da Placa Ativa"')],
    [TXT('    description = '), STR('"Liga e apaga o LED; adapta-se automaticamente ao BoardConfig"')],
    [''],
    [KW('    def '), DEF('_run'), TXT('(self, pal) -> TestResult:')],
    [KW('        from '), TXT('core.board.config '), KW('import '), TXT('get_active_board')],
    [TXT('        cfg  = get_active_board()')],
    [TXT('        port = cfg.led_port')],
    [TXT('        pin  = cfg.led_pin')],
    [''],
    [TXT('        pal.modify_field('), STR('"RCC"'), TXT(', cfg.led_clock_reg, cfg.led_clock_field, '), NUM('1'), TXT(')')],
    [TXT('        orig_moder = pal.read_field(port, '), STR('"MODER"'), TXT(', f'), STR('"MODER{pin}"'), TXT(')')],
    [''],
    [KW('        try'), TXT(':')],
    [TXT('            pal.modify_field(port, '), STR('"MODER"'), TXT(', f'), STR('"MODER{pin}"'), TXT(', '), NUM('0b01'), TXT(')')],
    [TXT('            pal.write_reg(port, '), STR('"BSRR"'), TXT(', '), NUM('1'), TXT(' << pin)            '), CMT('# SET')],
    [TXT('            alto  = pal.read_field(port, '), STR('"ODR"'), TXT(', f'), STR('"ODR{pin}"'), TXT(')')],
    [TXT('            pal.write_reg(port, '), STR('"BSRR"'), TXT(', '), NUM('1'), TXT(' << (pin + '), NUM('16'), TXT('))  '), CMT('# RESET')],
    [TXT('            baixo = pal.read_field(port, '), STR('"ODR"'), TXT(', f'), STR('"ODR{pin}"'), TXT(')')],
    [''],
    [KW('            if '), TXT('alto == '), NUM('1'), KW(' and '), TXT('baixo == '), NUM('0'), TXT(':')],
    [KW('                return '), TXT('self._pass('), STR('f"{port}{pin}: SET=1 RESET=0 OK"'), TXT(')')],
    [KW('            return '), TXT('self._fail(')],
    [TXT('                '), STR('f"SET={alto} RESET={baixo} (esperado 1 e 0)"'), TXT(',')],
    [TXT('                measured=(alto, baixo), expected=('), NUM('1'), TXT(', '), NUM('0'), TXT(')')],
    [TXT('            )')],
    [KW('        finally'), TXT(':')],
    [TXT('            pal.modify_field(port, '), STR('"MODER"'), TXT(', f'), STR('"MODER{pin}"'), TXT(', orig_moder)')],
], caption='LED portável usando get_active_board()')

h2('8.2 Exemplo: Verificar Configuração de USART pelo Firmware (Padrão B)')

para(
    'Objetivo: sem modificar nada, checar se o firmware já configurou a USART2 '
    'com 115200 baud, TX e RX habilitados.',
    space_after=6
)

code_block([
    [KW('from '), TXT('tests_builtin.base '), KW('import '), TXT('BuiltinTest, TestResult')],
    [KW('from '), TXT('instruments.usart '), KW('import '), TXT('UsartInstrument, _baud_divider')],
    [''],
    [KW('class '), DEF('Usart2ConfigCheck'), TXT('(BuiltinTest):')],
    [TXT('    name          = '), STR('"USART2 Config Check"')],
    [TXT('    BAUD_ESPERADO = '), NUM('115200')],
    [TXT('    F_CLK         = '), NUM('16_000_000')],
    [''],
    [KW('    def '), DEF('_run'), TXT('(self, pal) -> TestResult:')],
    [TXT('        u = UsartInstrument('), NUM('2'), TXT(', pal)')],
    [KW('        if not '), TXT('u.is_enabled:')],
    [KW('            return '), TXT('self._fail('), STR('"USART2 não habilitada (UE=0)"'), TXT(')')],
    [''],
    [TXT('        exp_m, exp_f = _baud_divider(self.BAUD_ESPERADO, self.F_CLK)')],
    [TXT('        got_m = pal.read_field('), STR('"USART2"'), TXT(', '), STR('"BRR"'), TXT(', '), STR('"DIV_Mantissa"'), TXT(')')],
    [TXT('        got_f = pal.read_field('), STR('"USART2"'), TXT(', '), STR('"BRR"'), TXT(', '), STR('"DIV_Fraction"'), TXT(')')],
    [''],
    [KW('        if '), TXT('got_m == exp_m '), KW('and '), TXT('got_f == exp_f:')],
    [KW('            return '), TXT('self._pass('), STR('f"USART2 OK: {self.BAUD_ESPERADO} baud"'), TXT(')')],
    [KW('        return '), TXT('self._fail(')],
    [TXT('            '), STR('f"BRR: mant={got_m}/{exp_m} frac={got_f}/{exp_f}"'), TXT(',')],
    [TXT('            measured=(got_m, got_f), expected=(exp_m, exp_f)')],
    [TXT('        )')],
], caption='USART config check — constantes fixas (Padrão B)')

quebra()


# ══════════════════════════════════════════════════════════════════════════════
# 9. COMO INTEGRAR
# ══════════════════════════════════════════════════════════════════════════════
h1('9. Como Integrar seu Teste na Plataforma')

para(
    'Há três formas de integrar um teste, dependendo de quantas vezes ele deve '
    'rodar e se faz parte de um produto específico ou da plataforma geral.',
    space_after=10
)

box(
    'Para testes parametrizáveis (blocos), o caminho recomendado é o Plano de testes do\n'
    'projeto (seção 14): o bloco é escrito uma vez e aplicado a pinos e periféricos\n'
    'escolhidos no formulário, sem editar código. As formas abaixo continuam valendo.',
    fill='E8F4F8',
    title='  Plano de testes',
    title_color=AZUL_MEDIO
)

h2('Forma 1 — Plugin via GUI (Recomendada para testes de produto)')

para(
    'Esta é a forma mais rápida e não requer nenhuma alteração no código da plataforma. '
    'Ideal quando o teste é específico para um produto ou cliente.',
    space_after=6
)

table(
    ['Passo', 'Ação'],
    [
        ['1', 'No painel esquerdo, selecione a placa no ComboBox "Projeto"'],
        ['2', 'Clique "Conectar" para conectar ao probe ST-Link'],
        ['3', 'Clique "Carregar Plugin..." no painel esquerdo'],
        ['4', 'Selecione o arquivo .py do seu teste'],
        ['5', 'O nome do teste aparece na lista "Plugins" no painel esquerdo'],
        ['6', 'Clique "Executar Suite" — seu teste roda junto com os embutidos'],
        ['7', 'Clique "Exportar PDF" para o relatório em PDF, "Relatório ▾" para HTML ou Word, ou "Exportar..." para os resultados em JSON ou CSV'],
        ['8', 'Escolha a pasta: a janela abre em Documentos na primeira vez e depois na última pasta usada; o nome sugerido leva projeto, data e hora'],
        ['9', 'Relatórios e JSON incluem o project_code e a seção "Instrumentos utilizados" (seção 13.3)'],
    ],
    widths=[1.5, 14.5],
    titulo='Integração de um teste como plugin via GUI',
)

figura('gui_aba_testes.png',
       'Figura 6 — Aba Testes: (1) testes embutidos; (2) "Carregar Plugin..."; '
       '(3) plugins carregados (aqui, GPIO Output)',
       largura_cm=6.5)

para(
    'Depois de carregar o plugin, clique "Executar Suite [F5]": o teste do plugin entra no '
    'fim da tabela de resultados, junto com os embutidos (veja a linha GPIO Output na Figura 2).',
    space_after=8
)

box(
    'O project_code aparece automaticamente no título da janela e no arquivo exportado.\n'
    'Isso permite saber exatamente qual configuração de placa gerou cada resultado.\n'
    'Exemplo: "Bring-up 2180.0720.00 — NUCLEO-F446RE"',
    fill='E8F4F8',
    title='  Rastreabilidade automática',
    title_color=AZUL_MEDIO
)

h2('Forma 2 — Suite Embutida (Para testes que sempre devem rodar)')

para(
    'Use esta forma quando o teste faz parte do processo padrão de verificação '
    'e deve rodar automaticamente em toda execução da plataforma.',
    space_after=6
)

code_block([
    [CMT('# src/tests_builtin/__init__.py')],
    [KW('from '), TXT('tests_builtin.meu_teste '), KW('import '), TXT('MeuTeste')],
    [''],
    [TXT('STANDARD_SUITE: tuple[BuiltinTest, ...] = (')],
    [TXT('    DeviceIdTest(),')],
    [TXT('    ClockTest(),')],
    [TXT('    SRAMTest(),')],
    [TXT('    LedTest(),')],
    [TXT('    ButtonTest(),')],
    [TXT('    UsartLoopbackTest(),')],
    [TXT('    SpiLoopbackTest(),')],
    [TXT('    MeuTeste(),    '), CMT('# <- seu teste aqui')],
    [TXT(')')],
    [''],
    [CMT('# Ou use build_suite(cfg) para testes unitários com placa explícita:')],
    [KW('from '), TXT('tests_builtin '), KW('import '), TXT('build_suite')],
    [KW('from '), TXT('boards.nucleo_f446re '), KW('import '), TXT('BOARD')],
    [TXT('suite = build_suite(BOARD)     '), CMT('# todos os testes com cfg injetado')],
], caption='src/tests_builtin/__init__.py')

h2('Forma 3 — CLI para Automação (CI/CD)')

code_block([
    [TXT('$env:PYTHONPATH = '), STR('"src"')],
    [TXT('.venv\\Scripts\\python.exe src/app/cli/main.py run --json > resultado.json')],
    [''],
    [CMT('# Verificar resultado (inclui project_code e board):')],
    [TXT('.venv\\Scripts\\python.exe -c "'), TXT('')],
    [TXT('import json, sys')],
    [TXT('data = json.load(open(\'resultado.json\'))')],
    [TXT('print(data[\'project_code\'], data[\'board\'])')],
    [TXT('falhas = [r for r in data[\'results\'] if r[\'status\'] != \'PASS\']')],
    [KW('if '), TXT('falhas:')],
    [TXT('    sys.exit('), NUM('1'), TXT(')')],
    [TXT('"')],
], caption='Automação via CLI com rastreamento de project_code')

quebra()


# ══════════════════════════════════════════════════════════════════════════════
# 10. PYTEST — GUIA COMPLETO
# ══════════════════════════════════════════════════════════════════════════════
h1('10. pytest — Guia Completo')

para(
    'pytest é o framework de testes unitários do projeto. Ele testa o CÓDIGO '
    'DA PLATAFORMA em si — não o hardware. Os testes em tests/ verificam se o '
    'PAL, os BuiltinTest, o runner e a GUI se comportam corretamente, usando '
    'FakeMemory (sem placa física).',
    space_after=8
)

h2('10.1 Metodologia — TDD')

para(
    'A metodologia seguida é Test-Driven Development (TDD). Antes de conectar a '
    'placa, o teste pytest deve: (1) passar com FakeMemory configurada com valores '
    'corretos e (2) falhar com valores errados. Se ambos funcionam, a lógica está '
    'certa. O hardware só entra depois.',
    space_after=6
)

code(
    '  Você escreve MeuTeste          Você escreve test_meu_teste.py\n'
    '  (BuiltinTest)                  (pytest)\n'
    '        |                               |\n'
    '        v                               v\n'
    '  Roda na GUI/CLI           Roda com: pytest tests/unit/ -v\n'
    '  com hardware real         com FakeMemory — sem hardware'
)

h2('10.2 Como o pytest descobre testes')

para(
    'O pytest segue convenções automáticas — nenhuma configuração extra é necessária:',
    space_after=4
)

table(
    ['Convenção', 'Exemplo'],
    [
        ['Arquivos começam com test_',       'test_builtins.py'],
        ['Classes começam com Test',          'class TestLed:'],
        ['Funções/métodos começam com test_', 'def test_pass_com_clock_ligado(pal):'],
    ],
    widths=[6, 10],
    titulo='Convenções de descoberta de testes do pytest',
)

h2('10.3 Fixtures — Injeção de Dependência')

para(
    'Fixtures preparam recursos para os testes (FakeMemory, PAL, SVD, conexão '
    'com hardware). O pytest injeta automaticamente pelo nome do parâmetro.',
    space_after=6
)

code_block([
    [TXT('@pytest.fixture')],
    [KW('def '), DEF('fake_mem'), TXT('():')],
    [KW('    return '), TXT('FakeMemory()              '), CMT('# criado antes de cada teste')],
    [''],
    [TXT('@pytest.fixture')],
    [KW('def '), DEF('pal'), TXT('(fake_mem, svd):   '), CMT('# injeta fake_mem e svd automaticamente')],
    [KW('    return '), TXT('PeripheralAccessLayer(fake_mem, svd)')],
    [''],
    [KW('def '), DEF('test_clock'), TXT('(pal):         '), CMT('# pytest injeta pal (que injeta fake_mem e svd)')],
    [TXT('    pal.modify_field('), STR('"RCC"'), TXT(', '), STR('"AHB1ENR"'), TXT(', '), STR('"GPIOAEN"'), TXT(', '), NUM('1'), TXT(')')],
    [TXT('    assert pal.read_field('), STR('"RCC"'), TXT(', '), STR('"AHB1ENR"'), TXT(', '), STR('"GPIOAEN"'), TXT(') == '), NUM('1')],
], caption='Fixture simples: injeção automática por nome')

h3('Escopos de fixture')

table(
    ['Escopo', 'Criado / destruido', 'Quando usar'],
    [
        ['function (padrão)', 'A cada teste',              'Estado que não deve vazar entre testes'],
        ['class',             'Uma vez por classe Test',   'Estado compartilhado entre métodos da classe'],
        ['module',            'Uma vez por arquivo .py',   'SVD parsing (lento, imutável)'],
        ['session',           'Uma vez por execução',      'Conexão USB com hardware real'],
    ],
    widths=[3.5, 5, 7.5],
    titulo='Escopos de fixture do pytest',
)

code_block([
    [TXT('@pytest.fixture(scope='), STR('"module"'), TXT(')    '), CMT('# SVD parseado 1x por arquivo de teste')],
    [KW('def '), DEF('svd'), TXT('():')],
    [KW('    return '), TXT('parse(Path('), STR('"devices/stm32f446/STM32F446.svd"'), TXT('))')],
    [''],
    [TXT('@pytest.fixture(scope='), STR('"session"'), TXT(')   '), CMT('# ST-Link aberto 1x por sessão inteira')],
    [KW('def '), DEF('stlink_adiv5'), TXT('(request):')],
    [TXT('    adiv5 = StLinkADIv5(target='), STR('"stm32f446retx"'), TXT(', mode='), STR('"halt"'), TXT(')')],
    [TXT('    adiv5._open()')],
    [KW('    yield '), TXT('adiv5          '), CMT('# yield: teardown acontece APÓS todos os testes')],
    [TXT('    adiv5._close()    '), CMT('# roda mesmo se os testes falharem')],
], caption='Escopos: module para SVD, session para hardware')

h3('yield — setup e teardown em uma fixture')

code_block([
    [TXT('@pytest.fixture')],
    [KW('def '), DEF('pal_com_clock'), TXT('(pal):')],
    [TXT('    pal.modify_field('), STR('"RCC"'), TXT(', '), STR('"AHB1ENR"'), TXT(', '), STR('"GPIOAEN"'), TXT(', '), NUM('1'), TXT(')  '), CMT('# setup')],
    [KW('    yield '), TXT('pal                                             '), CMT('# teste roda aqui')],
    [TXT('    pal.modify_field('), STR('"RCC"'), TXT(', '), STR('"AHB1ENR"'), TXT(', '), STR('"GPIOAEN"'), TXT(', '), NUM('0'), TXT(')  '), CMT('# teardown (sempre)')],
])

h3('autouse=True — fixture aplicada sem declarar')

code_block([
    [TXT('@pytest.fixture(autouse='), KW('True'), TXT(')')],
    [KW('def '), DEF('_active_board'), TXT('():')],
    [KW('    from '), TXT('boards.nucleo_f446re '), KW('import '), TXT('BOARD')],
    [TXT('    set_active_board(BOARD)   '), CMT('# roda antes de cada teste do módulo')],
    [CMT('                               # sem precisar declarar o parâmetro')],
])

h3('conftest.py — fixtures compartilhadas entre arquivos')

para(
    'O conftest.py é carregado automaticamente. Fixtures definidas nele ficam '
    'disponíveis para todos os arquivos da mesma pasta e subpastas, sem import:',
    space_after=4
)

code(
    'tests/\n'
    '├── conftest.py              <- fixtures disponíveis em tests/ inteiro\n'
    '│   (probe, probe_raw, probe_jtag)\n'
    '└── hardware/\n'
    '    ├── conftest.py          <- fixtures extras só para tests/hardware/\n'
    '    │   (stlink_adiv5, pal_stlink, hw_url, hw_freq)\n'
    '    └── test_pal_stlink.py  <- usa pal_stlink sem importar conftest'
)

h2('10.4 Marks — Metadados e Controle de Execução')

h3('@pytest.mark.parametrize — múltiplos cenários em um teste')

para(
    'Evita duplicar código para testar a mesma lógica com entradas diferentes. '
    'O pytest gera um teste separado (com nome próprio) para cada combinação:',
    space_after=6
)

code_block([
    [TXT('@pytest.mark.parametrize('), STR('"dev_id, status_esperado"'), TXT(', [')],
    [TXT('    ('), NUM('0x10000421'), TXT(', '), STR('"PASS"'), TXT('),    '), CMT('# F446 correto')],
    [TXT('    ('), NUM('0x10000123'), TXT(', '), STR('"FAIL"'), TXT('),    '), CMT('# chip errado')],
    [TXT('    ('), NUM('0x00000000'), TXT(', '), STR('"FAIL"'), TXT('),    '), CMT('# memória zerada')],
    [TXT('])')],
    [KW('def '), DEF('test_device_id'), TXT('(pal, fake_mem, dev_id, status_esperado):')],
    [TXT('    fake_mem._mem['), NUM('0xE0042000'), TXT('] = dev_id')],
    [TXT('    result = DeviceIdTest().run(pal)')],
    [TXT('    assert result.status.name == status_esperado')],
    [''],
    [CMT('# Saída:')],
    [CMT('# test_device_id[0x10000421-PASS]  PASSED')],
    [CMT('# test_device_id[0x10000123-FAIL]  PASSED')],
    [CMT('# test_device_id[0x00000000-FAIL]  PASSED')],
], caption='parametrize: 3 cenários, 3 testes separados')

h3('@pytest.mark.timeout — limite de tempo')

code_block([
    [CMT('# Instalar: .venv\\Scripts\\pip install pytest-timeout')],
    [''],
    [TXT('@pytest.mark.timeout('), NUM('5'), TXT(')   '), CMT('# falha se demorar mais de 5 segundos')],
    [KW('def '), DEF('test_loopback_nao_trava'), TXT('(pal_stlink):')],
    [TXT('    result = UsartLoopbackTest().run(pal_stlink)')],
    [TXT('    assert result.status != TestStatus.ERROR')],
    [''],
    [CMT('# Ou globalmente em pytest.ini:')],
    [CMT('# [pytest]')],
    [CMT('# timeout = 30    ; todos os testes têm limite de 30 s')],
])

h3('@pytest.mark.skip e skipif')

code_block([
    [TXT('@pytest.mark.skip(reason='), STR('"ainda não implementado"'), TXT(')')],
    [KW('def '), DEF('test_exportar_pdf'), TXT('():  ...')],
    [''],
    [TXT('@pytest.mark.skipif(sys.platform == '), STR('"win32"'), TXT(', reason='), STR('"só Linux"'), TXT(')')],
    [KW('def '), DEF('test_ftdi_url'), TXT('():  ...')],
    [''],
    [CMT('# Mark customizado (registrado em conftest.py):')],
    [TXT('@pytest.mark.hardware')],
    [KW('def '), DEF('test_led_real'), TXT('(pal_stlink):')],
    [TXT('    result = LedTest().run(pal_stlink)')],
    [TXT('    assert result.status == TestStatus.PASS')],
])

h2('10.5 Assertions')

code_block([
    [CMT('# assert puro — o mais comum:')],
    [TXT('assert result.status  == TestStatus.PASS')],
    [TXT('assert result.measured == '), NUM('0x421')],
    [TXT('assert '), STR('"1000"'), KW(' in '), TXT('result.message')],
    [''],
    [CMT('# pytest.raises — verificar exceção:')],
    [KW('def '), DEF('test_sem_placa'), TXT('():')],
    [TXT('    set_active_board('), KW('None'), TXT(')')],
    [KW('    with '), TXT('pytest.raises(RuntimeError, match='), STR('"Nenhuma placa"'), TXT('):')],
    [TXT('        get_active_board()')],
    [''],
    [CMT('# pytest.approx — floats com tolerância:')],
    [TXT('assert freq_medida == pytest.approx('), NUM('1000.0'), TXT(', rel='), NUM('0.01'), TXT(')')],
    [CMT('                                              # ±1% de tolerância')],
], caption='As três formas de assertion mais usadas')

h2('10.6 Testando código com threads — Filas')

para(
    'O TestRunnerService executa testes em uma thread separada e emite sinais Qt '
    'com os resultados. Para testar isso sem a GUI, use queue.Queue como ponte '
    'thread-safe: a thread do runner coloca resultados na fila (put) e a thread '
    'do teste retira (get) com timeout — sem race condition, sem sleep, sem polling.',
    space_after=6
)

code_block([
    [KW('import '), TXT('queue')],
    [''],
    [KW('def '), DEF('test_runner_emite_resultado'), TXT('(svd):')],
    [KW('    from '), TXT('app.services.runner '), KW('import '), TXT('TestRunnerService')],
    [KW('    from '), TXT('boards.nucleo_f446re '), KW('import '), TXT('BOARD')],
    [''],
    [TXT('    mem  = FakeMemory({'), NUM('0xE0042000'), TXT(': '), NUM('0x10000421'), TXT('})')],
    [TXT('    pal  = PeripheralAccessLayer(mem, svd)')],
    [TXT('    fila = queue.Queue()')],
    [''],
    [TXT('    svc = TestRunnerService()')],
    [TXT('    svc.result_ready.connect(fila.put)       '), CMT('# sinal Qt -> fila Python')],
    [TXT('    svc.start(pal, BOARD)')],
    [''],
    [TXT('    primeiro = fila.get(timeout='), NUM('5'), TXT(')           '), CMT('# aguarda até 5 s')],
    [TXT('    assert primeiro.name   == '), STR('"DeviceId"')],
    [TXT('    assert primeiro.status == TestStatus.PASS')],
], caption='queue.Queue como ponte entre thread do runner e thread do teste')

h2('10.7 Linha de comando')

code(
    '$env:PYTHONPATH = "src"\n'
    '\n'
    '# Rodar todos os testes unitários\n'
    '.venv\\Scripts\\python.exe -m pytest tests/unit/ -v\n'
    '\n'
    '# Só os testes cujo nome contém "led"\n'
    '.venv\\Scripts\\python.exe -m pytest tests/unit/ -k "led" -v\n'
    '\n'
    '# Parar no primeiro erro\n'
    '.venv\\Scripts\\python.exe -m pytest tests/unit/ -x\n'
    '\n'
    '# Mostrar os 5 testes mais lentos\n'
    '.venv\\Scripts\\python.exe -m pytest tests/unit/ --durations=5\n'
    '\n'
    '# Com hardware ST-Link (opção customizada do projeto)\n'
    '.venv\\Scripts\\python.exe -m pytest tests/hardware/ --stlink -v\n'
    '\n'
    '# Só testes sem hardware\n'
    '.venv\\Scripts\\python.exe -m pytest tests/ -m "not hardware"'
)

h2('10.8 As duas camadas no projeto')

table(
    ['Pasta', 'Usa hardware?', 'Scope típico', 'O que testa'],
    [
        ['tests/unit/',     'Não — FakeMemory',    'function / module', 'Lógica dos BuiltinTest, PAL, runner, GUI'],
        ['tests/hardware/', 'Sim — ST-Link / FT2232H', 'session',       'Comportamento real em registradores do STM32'],
    ],
    widths=[4, 4, 3.5, 4.5],
    titulo='As duas camadas de testes do projeto',
)

box(
    'Sem --stlink ou --hw-url, os testes de hardware são pulados automaticamente\n'
    'pelo conftest.py via pytest_collection_modifyitems. Isso permite rodar\n'
    'pytest tests/ -v em qualquer máquina sem hardware conectado.',
    fill='E8F4F8',
    title='  Skip automático sem hardware',
    title_color=AZUL_MEDIO
)

quebra()


# ══════════════════════════════════════════════════════════════════════════════
# 11. TDD
# ══════════════════════════════════════════════════════════════════════════════
h1('11. Desenvolvendo Sem Hardware (TDD)')

para(
    'Você não precisa ter a placa conectada para desenvolver e testar seu teste. '
    'A plataforma permite substituir o hardware real por uma memória falsa em Python '
    '(FakeMemory) que armazena os valores em um dicionário.',
    space_after=6
)
bullet('Desenvolver o teste no notebook, sem a placa')
bullet('Verificar se a lógica de pass/fail está correta')
bullet('Simular falhas (escrever valores errados e ver se o teste detecta)')
bullet('Rodar automaticamente em CI sem hardware')

doc.add_paragraph()
para('Crie o arquivo  tests/unit/test_meu_teste.py:')

code_block([
    [KW('from '), TXT('pathlib '), KW('import '), TXT('Path')],
    [KW('import '), TXT('pytest')],
    [KW('from '), TXT('core.board.config '), KW('import '), TXT('set_active_board')],
    [KW('from '), TXT('core.pal.access '), KW('import '), TXT('PeripheralAccessLayer')],
    [KW('from '), TXT('core.svd.parser '), KW('import '), TXT('parse')],
    [''],
    [KW('class '), DEF('FakeMemory'), TXT(':')],
    [KW('    def '), DEF('__init__'), TXT('(self): self._mem = {}')],
    [KW('    def '), DEF('read_mem32'), TXT('(self, addr): '), KW('return '), TXT('self._mem.get(addr & ~'), NUM('3'), TXT(', '), NUM('0'), TXT(')')],
    [KW('    def '), DEF('write_mem32'), TXT('(self, addr, value): self._mem[addr & ~'), NUM('3'), TXT('] = value & '), NUM('0xFFFFFFFF')],
    [''],
    [CMT('# Fixture obrigatória para testes que usam get_active_board()')],
    [TXT('@pytest.fixture(autouse='), KW('True'), TXT(')')],
    [KW('def '), DEF('board'), TXT('():')],
    [KW('    from '), TXT('boards.nucleo_f446re '), KW('import '), TXT('BOARD')],
    [TXT('    set_active_board(BOARD)')],
    [''],
    [TXT('@pytest.fixture')],
    [KW('def '), DEF('pal'), TXT('():')],
    [TXT('    svd = parse(Path('), STR('"devices/stm32f446/STM32F446.svd"'), TXT('))')],
    [KW('    return '), TXT('PeripheralAccessLayer(FakeMemory(), svd)')],
    [''],
    [KW('def '), DEF('test_led_passa'), TXT('(pal):')],
    [TXT('    pal.modify_field('), STR('"RCC"'), TXT(', '), STR('"AHB1ENR"'), TXT(', '), STR('"GPIOAEN"'), TXT(', '), NUM('1'), TXT(')')],
    [KW('    from '), TXT('meu_teste '), KW('import '), TXT('LedVerificacaoTest')],
    [KW('    from '), TXT('tests_builtin.base '), KW('import '), TXT('TestStatus')],
    [TXT('    result = LedVerificacaoTest().run(pal)')],
    [CMT('    # Pode passar ou falhar dependendo do FakeMemory')],
    [TXT('    assert result.status ', ), KW('in'), TXT(' (TestStatus.PASS, TestStatus.FAIL)')],
], caption='tests/unit/test_meu_teste.py')

para('Para rodar os testes unitários:')
code(
    '$env:PYTHONPATH = "src"\n'
    '.venv\\Scripts\\python.exe -m pytest tests/unit/ -v'
)

quebra()


# ══════════════════════════════════════════════════════════════════════════════
# 12. ALTA VELOCIDADE — AGENTE NA RAM
# ══════════════════════════════════════════════════════════════════════════════
h1('12. Testes de Alta Velocidade — Agente na RAM do MCU')

para(
    'Nos testes das seções anteriores, cada leitura ou escrita de registrador é uma ida e volta '
    'pelo USB até o probe. Medido com STLINK-V3 e NUCLEO-H743ZI: de 0,3 a 0,5 ms por acesso '
    '(cerca de 3 mil acessos por segundo) e de 120 a 400 KiB/s em leitura de bloco. Isso basta '
    'para configurar periféricos e conferir estados, mas não para alimentar um barramento rápido '
    'byte a byte: a 1 Mbit/s, um byte de USART dura 10 µs.',
    space_after=6
)
para(
    'A solução é tirar o SWD do caminho crítico. O host carrega um pequeno programa (stub) na '
    'RAM do MCU; o stub executa o teste na velocidade do núcleo e publica os resultados numa '
    'fila circular na RAM, que o host esvazia em paralelo e entrega a uma queue.Queue para análise.',
    space_after=8
)

h2('12.1 Quando usar cada caminho')
table(
    ['Situação', 'PAL (Python, seções 4 a 8)', 'Agente na RAM'],
    [
        ['Configuração, ID, estado de pinos, clocks', 'Sim', 'Desnecessário'],
        ['Loopback de 1 byte em qualquer baud rate', 'Sim — o byte trafega na velocidade real', 'Opcional'],
        ['Tráfego contínuo, vazão, overrun, BER', 'Não — o host cria buracos entre quadros', 'Sim'],
        ['Medir tempo com precisão de ciclo', 'Não', 'Sim — DWT->CYCCNT no próprio chip'],
        ['Firmware do produto continua rodando', 'Sim', 'Não — o stub assume o núcleo até o reset'],
    ],
    widths=[5.5, 5.0, 5.5],
    titulo='Quando usar o PAL ou o agente na RAM',
)

h2('12.2 Como funciona')
code(
    ' PC (Python)                                   STM32 (RAM)\n'
    ' 1. reset_and_halt + configura clock/GPIO      (clocks no estado de reset)\n'
    ' 2. carrega stub.bin e confere  ---- SWD ----> 0x24000000  código do stub\n'
    ' 3. preenche a caixa de correio ---- SWD ----> 0x24004000  parâmetros + fila\n'
    ' 4. PC = entrada, SP = topo, resume\n'
    '                                               stub roda a 64 MHz: varre os\n'
    '                                               baud rates, grava RX_BUF e um\n'
    '                                               registro por passo na fila\n'
    ' 5. RingReader.drain()  <------- SWD --------- head/tail da fila\n'
    '    -> queue.Queue -> thread de análise (analyze_ber...)\n'
    ' 6. stub termina em BKPT -> núcleo para -> host confere STATUS\n'
    ' 7. reset -> firmware do produto volta a rodar'
)

h2('12.3 Prova de conceito: velocidade máxima do USART (NUCLEO-H743ZI)')
para(
    'O USART3 é configurado em half-duplex: o receptor lê o próprio sinal pelo pino TX (PD8), '
    'sem fio externo. O kernel clock do USART é o HSI de 64 MHz, o que limita a varredura a '
    '8 Mbaud (f/8 com OVER8).',
    space_after=6
)
code(
    'python tools/agent_usart_sweep.py --bytes 4096',
    caption='Rodar a partir da raiz do repositório, com a placa ligada ao STLINK-V3'
)
code(
    '   baud     real    rx/tx    erros ORE FE NE  BER       efetivo     efic.\n'
    ' 115200   115108  4096/4096   0    0   0  0  0.00e+00  0.115 Mb/s  100%  OK\n'
    '1000000  1000000  4096/4096   0    0   0  0  0.00e+00  1.000 Mb/s  100%  OK\n'
    '4000000  4000000  4096/4096   0    0   0  0  0.00e+00  3.999 Mb/s  100%  OK\n'
    '6400000  6400000  4096/4096   0    0   0  0  0.00e+00  6.398 Mb/s  100%  OK\n'
    '8000000  8000000  4096/4096   0    0   0  0  0.00e+00  6.725 Mb/s   84%  OK\n'
    'Maior baud sem erro: 8000000 (real 8000000), taxa efetiva 6.725 Mb/s em 6.09 ms',
    caption='Saída na bancada (colunas condensadas, linhas intermediárias omitidas)'
)
bullet('real — baud rate obtido com o divisor inteiro do BRR (ex.: 921600 vira 927536)')
bullet('erros / ORE / FE / NE — bytes diferentes do padrão, overrun, erro de quadro e ruído')
bullet('BER — calculado pelo analyze_ber comparando RX_BUF com o padrão enviado')
bullet('efetivo — bits no fio (10 por byte) divididos pelo tempo medido no DWT')
bullet('efic. abaixo de 100% — o stub não alimentou a FIFO a tempo: limite do programa '
       '(64 MHz, código na AXI SRAM), não do barramento')

h2('12.4 Usando o agente no seu código')
code_block([
    [KW('import '), TXT('queue, threading')],
    [KW('from '), TXT('adapters.stlink.adiv5 '), KW('import '), TXT('StLinkADIv5')],
    [KW('from '), TXT('core.agent '), KW('import '), TXT('run_sweep, max_reliable_baud')],
    '',
    [TXT('q = queue.Queue()')],
    [KW('def '), DEF('analisar'), TXT('():')],
    [KW('    while '), TXT('(r := q.get()) '), KW('is not None'), TXT(':     '),
     CMT('# chega enquanto o stub roda')],
    [TXT('        print(r.baud_nominal, r.ok, r.ber.value)')],
    [TXT('threading.Thread(target=analisar, daemon='), KW('True'), TXT(').start()')],
    '',
    [KW('with '), TXT('StLinkADIv5(target='), STR('"stm32h743xx"'), TXT(', mode='),
     STR('"under-reset"'), TXT(') '), KW('as '), TXT('t:')],
    [TXT('    t.reset_and_halt()                  '), CMT('# clocks no estado de reset')],
    [TXT('    res = run_sweep(t, n_bytes='), NUM('1024'), TXT(', out=q)')],
    [TXT('    t.reset_system()                    '), CMT('# firmware volta a rodar')],
    '',
    [TXT('melhor = max_reliable_baud(res)        '), CMT('# maior baud sem erro')],
], caption='Varredura com análise em paralelo pela fila')

h2('12.5 Criando um novo stub')
table(
    ['Passo', 'Ação'],
    [
        ['1', 'Copie stubs/stm32h7/usart_sweep.c e stub.ld; mantenha a função de entrada na seção .entry'],
        ['2', 'Não use variáveis globais (.data/.bss são descartados: não há código de startup)'],
        ['3', 'Publique cada registro na fila, execute DSB e só então incremente head'],
        ['4', 'Termine com STATUS = DONE e BKPT — o host detecta o núcleo parado'],
        ['5', 'Registre o stub em STUBS, no stubs/build_stubs.py, e rode python stubs/build_stubs.py'],
        ['6', 'Escreva o lado host (configuração via SWD + caixa de correio) usando TargetAgent e RingReader'],
        ['7', 'Teste o lado host sem placa com um alvo simulado (ver tests/unit/test_agent.py)'],
    ],
    widths=[1.5, 14.5],
    titulo='Criação de um novo stub de RAM',
)

table(
    ['Endereço (STM32H7)', 'Conteúdo'],
    [
        ['0x24000000 (AXI SRAM)', 'Código do stub (até 16 KiB)'],
        ['0x24004000', 'Caixa de correio: MAGIC, STATUS, parâmetros, head/tail, passos e fila'],
        ['0x24010000', 'RX_BUF: bytes recebidos, n bytes por passo'],
        ['0x20020000 (topo da DTCM)', 'Pilha do stub'],
    ],
    widths=[5.0, 11.0],
    titulo='Mapa de memória do stub (STM32H7)',
)

box(
    'O stub assume o núcleo: o firmware do produto para durante o teste e volta com reset.\n'
    'Placas cujo firmware entra em baixo consumo só aceitam conexão com mode="under-reset".\n'
    'PD8 também vai ao Virtual COM do ST-Link: um terminal aberto mostra lixo no teste.\n'
    'Acima de 8 Mbaud é preciso ligar o PLL; no H7 isso exige configurar PWR_CR3, que só\n'
    'aceita uma escrita após o power-up — valor errado pode impedir o chip de subir.\n'
    'O .bin é versionado; registre o sha256 impresso pelo build_stubs.py junto ao resultado.\n'
    'Endereços e bits seguem o RM0433 e devem ser revisados antes de virar critério de aprovação.',
    fill=FUNDO_AVISO,
    title='  Cuidados',
    title_color=VERMELHO
)

quebra()


# ══════════════════════════════════════════════════════════════════════════════
# 13. INSTRUMENTOS DE MEDIÇÃO
# ══════════════════════════════════════════════════════════════════════════════
h1('13. Instrumentos de Medição')

para(
    'Além de ler registradores do MCU, um teste pode usar instrumentos de bancada — '
    'osciloscópio, analisador de espectro, multímetro — para medir o que acontece no fio. '
    'A plataforma separa três camadas, no mesmo espírito do BoardConfig:',
    space_after=6
)
table(
    ['Camada', 'O que faz', 'Onde fica'],
    [
        ['Transporte', 'Fala com o aparelho: VISA (LAN, USB-TMC, GPIB), socket SCPI, serial, mock',
         'src/instruments/drivers/'],
        ['Função', 'Interface comum por tipo: Oscilloscope, SpectrumAnalyzer',
         'src/core/instruments/oscilloscope.py, spectrum.py'],
        ['Modelo', 'Comandos SCPI de cada fabricante, escolhido pelo *IDN?',
         'src/instruments/models/'],
    ],
    widths=[2.8, 8.2, 5.0],
    titulo='Camadas da integração de instrumentos',
)
para(
    'O teste pede a função ("o osciloscópio"), nunca o modelo. Trocar o Rigol por outro '
    'osciloscópio suportado não muda nenhuma linha do teste.',
    space_after=8
)

h2('13.1 Modelos suportados')
table(
    ['Modelo', 'Função', 'Situação'],
    [
        ['Rigol MSO8000 (MSO8064/8104/8204)', 'Osciloscópio',
         'Implementado; comandos ainda não validados em bancada'],
        ['Rohde & Schwarz FPC1000/FPC1500', 'Analisador de espectro',
         'Conceitual: escrito pela sintaxe SCPI da família, sem acesso ao aparelho'],
        ['Qualquer aparelho SCPI', '— (SCPI direto)', 'get_instrument("nome").query("...")'],
    ],
    widths=[5.5, 3.5, 7.0],
    titulo='Modelos de instrumento suportados',
)

h2('13.2 Cadastrar e conectar')
para(
    'Na aba Instrumentos, clique "Gerenciar...". O caminho mais rápido é "Detectar VISA...": '
    'a plataforma lista os recursos VISA da rede e das portas USB/GPIB, pergunta *IDN? a cada um '
    'e mostra o resultado para você marcar o que cadastrar. Portas seriais ficam de fora '
    '(costumam ser a porta COM virtual do ST-Link) e o mesmo aparelho visto por dois caminhos '
    '(VXI-11 e HiSLIP) aparece uma vez só.',
    space_after=6
)
figura('gui_instr_detectar.png',
       'Figura 7 — "Detectar VISA...": aparelhos que responderam ao *IDN? (dados simulados)',
       largura_cm=13)
para(
    'Os marcados entram na lista com nome sugerido — "osc" para osciloscópio, "sa" para '
    'analisador de espectro — que é o nome que os testes usam. Selecione e clique "Conectar": '
    'a plataforma lê o *IDN? de novo e guarda a identificação para o relatório.',
    space_after=6
)
figura('gui_instr_gerenciar.png',
       'Figura 8 — Gerenciador: (1) "Detectar VISA..."; (2) instrumentos cadastrados. A bolinha '
       'fica laranja quando a calibração vence em até 30 dias e vermelha quando já venceu')
box(
    'Prefira a conexão por rede (LAN). Ela funciona com o PyVISA puro (pyvisa-py), já instalado.\n'
    'Pela USB, o Windows precisa de uma biblioteca VISA do fabricante (NI-VISA ou equivalente).\n'
    'Instalação: pip install -e ".[hardware]" (inclui pyvisa e pyvisa-py).',
    fill='E8F4F8',
    title='  Rede ou USB?',
    title_color=AZUL_MEDIO
)

h2('13.3 Rastreabilidade e calibração')
para(
    'Cada relatório (HTML, Word e PDF) e a exportação JSON trazem a seção "Instrumentos '
    'utilizados", com fabricante, modelo, número de série, firmware, patrimônio e validade de '
    'calibração de cada aparelho conectado na sessão.',
    space_after=6
)
table(
    ['Campo no cadastro', 'Chave em instruments.json', 'Efeito'],
    [
        ['Patrimônio', 'asset_id', 'Aparece no relatório'],
        ['Calibração até (AAAA-MM-DD)', 'calibration_due',
         'Relatório + cor no gerenciador + aviso no log ao conectar'],
        ['—', 'block_if_expired: true', 'Recusa conectar com calibração vencida'],
        ['—', 'model: "rigol_mso8000"', 'Força o modelo quando o *IDN? não é reconhecido'],
    ],
    widths=[4.5, 4.8, 6.7],
    titulo='Campos de rastreabilidade no cadastro do instrumento',
)

h2('13.4 Usando nos testes')
code_block([
    [KW('from '), TXT('core.analysis '), KW('import '), TXT('MeasurementSet, analyze_frequency')],
    [KW('from '), TXT('core.instruments.roles '), KW('import '), TXT('get_oscilloscope')],
    '',
    [KW('def '), DEF('_run'), TXT('(self, pal) -> TestResult:')],
    [TXT('    osc = get_oscilloscope('), STR('"osc"'), TXT(')')],
    [KW('    if '), TXT('osc '), KW('is None'), TXT(':          '),
     CMT('# sem aparelho: SKIP, não FAIL')],
    [KW('        return '), TXT('self._skip('), STR('"Osciloscópio não conectado"'), TXT(')')],
    [TXT('    osc.configure_channel('), NUM('1'), TXT(', scale_v='), NUM('1.0'), TXT(')')],
    [TXT('    osc.set_timebase('), NUM('100e-9'), TXT(')')],
    [TXT('    osc.set_edge_trigger('), NUM('1'), TXT(', level_v='), NUM('1.65'), TXT(')')],
    [TXT('    osc.acquire_single(timeout_s='), NUM('3'), TXT(')        '),
     CMT('# arma e espera o disparo')],
    [TXT('    wf = osc.get_waveform('), NUM('1'), TXT(')                  '),
     CMT('# volts + dt (s)')],
    '',
    [TXT('    ms = MeasurementSet()')],
    [TXT('    ms.add(analyze_frequency('), STR('"f"'), TXT(', wf.edges('), NUM('1.65'),
     TXT('), '), NUM('4e6'), TXT('))')],
    [TXT('    f_auto = osc.measure('), STR('"frequency"'), TXT(', '), NUM('1'), TXT(')  '),
     CMT('# medida do próprio aparelho')],
], caption='Osciloscópio dentro de um teste')
para(
    'wf.edges(nível) devolve os instantes de cada borda, com interpolação linear entre amostras. '
    'Essa lista alimenta diretamente analyze_frequency, analyze_timing e analyze_jitter, com os '
    'mesmos critérios de aprovação dos demais testes. get_waveform(ch, raw=True) lê a memória '
    'completa do osciloscópio, em blocos (mais lento).',
    space_after=6
)
table(
    ['Osciloscópio', 'Analisador de espectro'],
    [
        ['configure_channel, set_timebase, set_edge_trigger', 'set_center_span, set_start_stop'],
        ['single, acquire_single, run, stop, force_trigger', 'set_rbw, set_ref_level'],
        ['measure(item, ch) — frequency, vpp, rise_time, duty...', 'single_sweep'],
        ['get_waveform(ch) → Waveform (.edges, .times)', 'get_trace() → Spectrum (.peak, .level_at)'],
        ['', 'peak() — marcador no pico'],
    ],
    widths=[8.5, 7.5],
    titulo='Métodos do osciloscópio e do analisador de espectro',
)
para(
    'Exemplo completo: TESTES_PLATAFORMA/INSTRUMENTOS/teste_osc_mco.py. O MCU coloca o HSI/4 '
    '(4 MHz) no pino PA8 (MCO1) da NUCLEO-F446RE e o osciloscópio mede a frequência pelas '
    'bordas e pela medida automática; os registradores alterados são restaurados no finally.',
    space_after=8
)

h2('13.5 Incluindo um modelo novo')
table(
    ['Passo', 'Ação'],
    [
        ['1', 'Crie src/instruments/models/<fabricante_modelo>.py herdando de Oscilloscope ou '
              'SpectrumAnalyzer'],
        ['2', 'Implemente cada método com os comandos do manual de programação do aparelho'],
        ['3', 'Acrescente uma linha em MODELS (instruments/models/__init__.py) com as expressões '
              'regulares de fabricante e modelo do *IDN?'],
        ['4', 'Teste com o MockDriver: confira a sequência exata de comandos e a conversão de '
              'dados (ver tests/unit/test_instrument_classes.py)'],
        ['5', 'Valide na bancada e marque validated=True na linha do MODELS'],
    ],
    widths=[1.5, 14.5],
    titulo='Inclusão de um modelo de instrumento',
)
box(
    'Os comandos SCPI do Rigol MSO8000 e do R&S FPC foram escritos a partir dos manuais de\n'
    'programação, sem o aparelho na bancada. Antes de servirem de critério de aprovação, cada\n'
    'comando deve ser conferido com o manual da versão de firmware e com o equipamento, e o\n'
    'resultado revisado por pessoa qualificada distinta de quem escreveu o driver.',
    fill=FUNDO_AVISO,
    title='  Validação pendente',
    title_color=VERMELHO
)

quebra()


# ══════════════════════════════════════════════════════════════════════════════
# 14. PLANO DE TESTES E BLOCOS
# ══════════════════════════════════════════════════════════════════════════════
h1('14. Plano de Testes e Blocos Reutilizáveis')

para(
    'Um bloco é um teste escrito uma vez, sem nenhum pino ou registrador fixo: ele declara do '
    'que precisa (pinos, um periférico, números, opções) e o projeto diz a que se aplica. A '
    'suite do produto é montada combinando blocos no Plano de testes do projeto.',
    space_after=6
)

h2('14.1 Montando o plano')
para(
    'Na aba Testes da janela principal, clique "Editar plano de testes..." (ou dê dois cliques '
    'na lista do plano). O editor do projeto abre na aba "Plano de testes".',
    space_after=6
)
figura('gui_plano_testes.png',
       'Plano de testes: (1) instâncias, marcadas = executadas, na ordem da lista; (2) formulário '
       'gerado pelos parâmetros do bloco; (3) "+ Adicionar teste..."')
table(
    ['Passo', 'Ação'],
    [
        ['1', 'Clique "+ Adicionar teste..." e escolha o bloco na biblioteca (agrupada por categoria)'],
        ['2', 'Dê um rótulo: ele entra no nome do resultado ("UART — Loopback — USART3 VCP 115200")'],
        ['3', 'Escolha os pinos (lista do Pinout, filtrada pelo tipo) ou o periférico (USART3, I2C1...)'],
        ['4', 'Ajuste os demais parâmetros e repita para outras instâncias; Duplicar ajuda a variar'],
        ['5', 'Salve: o plano fica no JSON do projeto e roda com "Executar Suite", após os embutidos'],
    ],
    widths=[1.5, 14.5],
    titulo='Montagem do plano de testes',
)
figura('gui_adicionar_teste.png',
       'Biblioteca de blocos agrupada por categoria; "Outro arquivo..." aceita blocos de qualquer '
       'pasta', largura_cm=12)
para(
    '"Carregar Plugin..." também leva ao plano: se o arquivo tiver um teste com parâmetros, a '
    'plataforma oferece adicioná-lo ao plano do projeto e já abre o formulário.',
    space_after=8
)

h2('14.2 Blocos disponíveis')
table(
    ['Categoria', 'Bloco', 'O que valida'],
    [
        ['GPIO', 'Piscar LED', 'Saídas acendem e apagam (ODR conferido); respeita ativo-baixo'],
        ['GPIO', 'Nível de entrada', 'Entradas no nível esperado (ativo, inativo ou só registra)'],
        ['GPIO', 'Pull-up/pull-down (curto)', 'Pino segue o pull interno; acusa curto a GND/VCC'],
        ['GPIO', 'Loopback saída→entrada', 'Uma saída aciona uma entrada (fio ou circuito)'],
        ['Interrupções', 'EXTI de um pino', 'Pino → EXTI → pendência no EXTI e no NVIC'],
        ['UART', 'Loopback', 'N bytes com BER; half-duplex sem fio ou fio TX→RX'],
        ['SPI', 'Loopback MOSI→MISO', 'Palavras com BER (jumper MOSI→MISO)'],
        ['I2C', 'Varredura do barramento', 'Endereços que respondem; esperados; barramento preso'],
        ['CAN', 'Loopback interno', 'Quadros TX→RX no controlador, contadores de erro'],
        ['Ethernet', 'PHY por MDIO', 'ID do PHY e link (cabo), pelo MDIO do MAC'],
    ],
    widths=[2.6, 4.6, 8.8],
    titulo='Blocos de TESTES_PLATAFORMA (BLOCOS e IO)',
)
para(
    'Os blocos de GPIO são montados com o testbench modular (seção 15): estímulo, drivers, '
    'monitores e decisões reutilizáveis.',
    space_after=6
)
table(
    ['Bloco', 'Resultado'],
    [
        ['UART half-duplex, USART3', 'PASS a 115200 e a 2 Mbaud, BER 0'],
        ['I2C1, varredura', 'PASS: barramento livre, nenhum dispositivo na Nucleo'],
        ['SPI1, loopback', 'Transfere; sem o jumper D11→D12, MISO fica em 0 (FAIL esperado)'],
        ['FDCAN1, loopback interno', 'PASS: 8 quadros a 500 kbit/s, TEC=0, REC=0'],
        ['Ethernet, PHY 0', 'PASS: LAN8742A identificado (0007:C131), sem link sem cabo'],
        ['EXTI13 (botão B1)', 'PASS por software e por borda real no pino'],
    ],
    widths=[5.5, 10.5],
    titulo='Validação dos blocos na NUCLEO-H753ZI',
)
box(
    'O caminho F4 dos blocos (USART v1, SPI v1, I2C v1, bxCAN) segue o RM0390 mas ainda não\n'
    'foi executado em placa. Os bits de todos os blocos seguem RM0390/RM0433 e precisam de\n'
    'revisão técnica por outra pessoa antes de servirem de critério de aprovação.',
    fill=FUNDO_AVISO,
    title='  Validação pendente',
    title_color=VERMELHO
)

h2('14.3 A camada de periféricos (src/hal)')
para(
    'Os blocos não acessam registradores de comunicação diretamente: usam a HAL, que escolhe a '
    'implementação pela geração do periférico detectada no SVD e restaura no fim os '
    'registradores e pinos que alterou.',
    space_after=6
)
table(
    ['Módulo', 'Gerações', 'Observações'],
    [
        ['hal.common', '—', 'Família, bit de clock achado no RCC do SVD, AF padrão, salvar/restaurar pino'],
        ['hal.uart', 'v1 (F4: SR/DR), v2 (H7: ISR/TDR/RDR)', 'H7: kernel no HSI durante o teste'],
        ['hal.spi', 'v1 (F4), v2 (H7: CFG1/CFG2/TXDR)', 'H7: SCK precisa estar em AF para gerar clock'],
        ['hal.i2c', 'v1 (F4: SR1/SR2), v2 (H7: TIMINGR/ISR)', 'TIMINGR calculado para ~100 kHz'],
        ['hal.exti', 'IMR/PR (F4), CPUIMR1/CPUPR1 (H7)', 'Confere também a pendência no NVIC'],
        ['hal.can', 'bxCAN (F4), FDCAN (H7)', 'SKIP se o firmware estiver usando o CAN'],
        ['hal.eth', 'MACMIIAR (F4), MACMDIOAR (H7)', 'Só leitura de registradores do PHY'],
    ],
    widths=[2.6, 6.0, 7.4],
    titulo='Módulos da HAL de periféricos',
)
para(
    'A função alternativa (AF) de cada pino vem de uma tabela por família. Para um pino com AF '
    'fora do padrão, informe-a no próprio pino: PeripheralRef.params["af"].',
    space_after=8
)

h2('14.4 Escrevendo um bloco')
code_block([
    [KW('from '), TXT('tests_builtin.base '), KW('import '), TXT('BuiltinTest, TestResult')],
    [KW('from '), TXT('tests_builtin.params '), KW('import '), TXT('Int, Periph, Signals')],
    '',
    [KW('class '), DEF('MeuBloco'), TXT('(BuiltinTest):')],
    [TXT('    name        = '), STR('"UART — Eco simples"')],
    [TXT('    category    = '), STR('"UART"'), TXT('              '),
     CMT('# grupo na biblioteca')],
    [TXT('    description = '), STR('"Envia um byte e espera o eco"')],
    [TXT('    params = {')],
    [TXT('        '), STR('"periferico"'), TXT(': Periph('), STR('"Periférico"'), TXT(', ('),
     STR('"UART"'), TXT(',)),')],
    [TXT('        '), STR('"baud"'), TXT(':       Int('), STR('"Baud rate"'), TXT(', '),
     NUM('115200'), TXT(', minimum='), NUM('300'), TXT(', maximum='), NUM('12_000_000'),
     TXT('),')],
    [TXT('        '), STR('"led"'), TXT(':        Signals('), STR('"LED de status"'),
     TXT(', max_count='), NUM('1'), TXT('),')],
    [TXT('    }')],
    '',
    [KW('    def '), DEF('_run'), TXT('(self, pal) -> TestResult:')],
    [TXT('        tx = self.p.periferico.pin('), STR('"TX"'), TXT(')   '),
     CMT('# Signal: .port .pin .af')],
    [TXT('        led = self.p.led['), NUM('0'), TXT(']')],
    [TXT('        ...                                  '), CMT('# use hal.uart.Uart')],
], caption='Bloco com parâmetros: aparece sozinho na biblioteca')
table(
    ['Parâmetro', 'No formulário', 'No teste (self.p.<nome>)'],
    [
        ['Signals(...)', 'Sinais do Pinout com caixas de seleção', 'lista de Signal (name, port, pin...)'],
        ['Periph(...)', 'Lista de periféricos do tipo', 'Peripheral: .pin("TX"), .has("RX")'],
        ['Int / Float', 'Campo numérico com limites', 'int / float'],
        ['Choice', 'Lista de opções', 'str'],
        ['Bool', 'Caixa de seleção', 'bool'],
        ['Text', 'Campo de texto', 'str'],
    ],
    widths=[3.0, 6.0, 7.0],
    titulo='Tipos de parâmetro (tests_builtin.params)',
)
para(
    'Cada resultado registra os parâmetros usados (relatório e JSON). Parâmetros inválidos — '
    'sinal que não existe mais no Pinout, valor fora dos limites — viram ERROR com a lista dos '
    'problemas, sem executar o bloco.',
    space_after=8
)

quebra()


# ══════════════════════════════════════════════════════════════════════════════
# 15. TESTBENCH MODULAR
# ══════════════════════════════════════════════════════════════════════════════
h1('15. Testbench Modular: Estímulo, Drivers, Monitores e Decisões')

para(
    'Um teste escrito "de uma vez" mistura quatro coisas: o que fazer, como agir no hardware, '
    'como observar e como julgar. O testbench (src/testbench) separa esses papéis em peças '
    'reutilizáveis. O mesmo driver de saída serve ao Piscar LED e ao loopback; o mesmo monitor '
    'de entrada serve à leitura de nível e ao teste de curto; a decisão "acompanhou o valor '
    'aplicado?" serve a qualquer par driver/monitor.',
    space_after=6
)
figura('diagrama_testbench.png',
       'Fluxo de dados de um teste composto: estímulo → drivers → hardware → monitores → '
       'histórico → decisões → veredito')

h2('15.1 Os quatro papéis')
table(
    ['Papel', 'Pergunta que responde', 'Componentes prontos'],
    [
        ['Estímulo', 'O que fazer, em que ordem e com que espera?',
         'Toggle (liga/desliga N vezes), Levels ([0, 1, 0]), Sample (só lê), Chain (concatena)'],
        ['Driver', 'Como aplicar o valor do passo no hardware?',
         'GpioOutputDriver (saída, respeita ativo-baixo), GpioPullDriver (pull-up/down)'],
        ['Monitor', 'O que observar a cada passo?',
         'GpioInputMonitor (IDR), GpioLatchMonitor (ODR)'],
        ['Decisão', 'O observado é o esperado?',
         'FollowsDrive (acompanhou o driver), Expect (valor esperado), Report (só registra)'],
    ],
    widths=[2.4, 5.0, 8.6],
    titulo='Papéis do testbench e componentes disponíveis',
)
para(
    'O estímulo não conhece hardware: ele só diz "driver X recebe o valor V e espera T". Drivers '
    'e monitores recebem o sinal do Pinout (nome, porta, pino, ativo-baixo) e salvam o estado '
    'do pino no setup() para restaurá-lo no teardown(). Decisões só leem o histórico (Trace).',
    space_after=8
)

h2('15.2 Ciclo de execução')
figura('diagrama_ciclo.png',
       'Testbench.run(): setup → passos (aplica, espera, amostra, grava) → teardown → decisões '
       '→ resultado; o teardown acontece mesmo se houver erro no meio')
table(
    ['Etapa', 'O que acontece'],
    [
        ['1. setup', 'Drivers e depois monitores: salvam o estado do pino e o configuram'],
        ['2. passos', 'Para cada Step: aplica o valor em cada driver, espera wait_s, amostra os '
                      'monitores e grava uma Observation (aplicado + observado) no Trace'],
        ['3. teardown', 'Monitores e depois drivers (ordem inversa): restauram os pinos — sempre, '
                        'mesmo com exceção (finally)'],
        ['4. decisões', 'Cada decisão avalia o Trace e devolve um Verdict: ok, falha ou '
                        'informativo, com métricas'],
        ['5. resultado', 'As métricas entram num MeasurementSet (o perfil de critérios é '
                         'aplicado) e o teste vira PASS ou FAIL'],
    ],
    widths=[2.6, 13.4],
    titulo='Etapas de Testbench.run()',
)

h2('15.3 Exemplo: o Piscar LED montado')
code_block([
    [KW('from '), TXT('testbench '), KW('import '), TXT('ComposedTest, Testbench')],
    [KW('from '), TXT('testbench.decision '), KW('import '), TXT('FollowsDrive')],
    [KW('from '), TXT('testbench.gpio '), KW('import '), TXT('GpioLatchMonitor, GpioOutputDriver')],
    [KW('from '), TXT('testbench.stimulus '), KW('import '), TXT('Toggle')],
    '',
    [KW('class '), DEF('PiscarLed'), TXT('(ComposedTest):')],
    [TXT('    params = { ... }                      '), CMT('# pinos, ciclos, periodo_ms, sequencial')],
    '',
    [KW('    def '), DEF('build'), TXT('(self) -> Testbench:')],
    [TXT('        drivers  = [GpioOutputDriver(s, initial='), KW('False'),
     TXT(') '), KW('for '), TXT('s '), KW('in '), TXT('self.p.pinos]')],
    [TXT('        monitors = [GpioLatchMonitor(s) '), KW('for '), TXT('s '), KW('in '),
     TXT('self.p.pinos]')],
    [KW('        return '), TXT('Testbench(')],
    [TXT('            drivers=drivers, monitors=monitors,')],
    [TXT('            stimulus=Toggle([d.name '), KW('for '), TXT('d '), KW('in '),
     TXT('drivers], self.p.ciclos,')],
    [TXT('                            self.p.periodo_ms / '), NUM('1000'),
     TXT(', self.p.sequencial),')],
    [TXT('            decisions=[FollowsDrive(m.name, d.name, label=s.name)')],
    [TXT('                       '), KW('for '), TXT('d, m, s '), KW('in '),
     TXT('zip(drivers, monitors, self.p.pinos)]),')],
], caption='TESTES_PLATAFORMA/IO/piscar_led.py — só monta a bancada; executar e restaurar é do framework')
table(
    ['Script', 'Estímulo', 'Drivers', 'Monitores', 'Decisão'],
    [
        ['Piscar LED', 'Toggle', 'GpioOutputDriver', 'GpioLatchMonitor', 'FollowsDrive'],
        ['GPIO Output', 'Levels [1, 0]', 'GpioOutputDriver', 'GpioLatchMonitor', 'FollowsDrive'],
        ['GPIO Input', 'Sample', '—', 'GpioInputMonitor', 'Report'],
        ['Nível de entrada', 'Sample', '—', 'GpioInputMonitor', 'Expect ou Report'],
        ['Pull (curto)', 'pull-up, pull-down', 'GpioPullDriver', 'GpioInputMonitor', 'Expect'],
        ['Loopback GPIO', 'Levels [0, 1, 0]', 'GpioOutputDriver', 'GpioInputMonitor', 'FollowsDrive'],
    ],
    widths=[3.0, 3.0, 3.4, 3.4, 3.2],
    titulo='Scripts de GPIO montados com o testbench',
)
para(
    'Esses scripts mantiveram nome de classe, arquivo e parâmetros: planos de teste salvos antes '
    'da mudança continuam válidos. Os blocos de comunicação (UART, SPI, I2C, CAN, Ethernet, EXTI) '
    'ainda seguem o formato anterior e podem ser convertidos com drivers e monitores próprios '
    '(ex.: driver de TX e monitor de RX de UART, decisão de BER).',
    space_after=8
)

h2('15.4 Escrevendo um driver, um monitor ou uma decisão')
code_block([
    [KW('from '), TXT('testbench '), KW('import '), TXT('Decision, Driver, Monitor, Verdict')],
    '',
    [KW('class '), DEF('DacDriver'), TXT('(Driver):            '), CMT('# aplica tensão no DAC')],
    [KW('    def '), DEF('__init__'), TXT('(self, canal): self.name = '), STR('f"dac:{canal}"'),
     TXT('; self.canal = canal')],
    [KW('    def '), DEF('setup'), TXT('(self, pal): ...       '), CMT('# salva estado, liga o DAC')],
    [KW('    def '), DEF('apply'), TXT('(self, pal, volts): ...'), CMT('# escreve o código no DHR12R')],
    [KW('    def '), DEF('teardown'), TXT('(self, pal): ...    '), CMT('# restaura')],
    '',
    [KW('class '), DEF('AdcMonitor'), TXT('(Monitor):          '), CMT('# lê o ADC a cada passo')],
    [KW('    def '), DEF('sample'), TXT('(self, pal) -> float: ...')],
    '',
    [KW('class '), DEF('DentroDaFaixa'), TXT('(Decision):       '), CMT('# decide pelo Trace')],
    [KW('    def '), DEF('evaluate'), TXT('(self, trace) -> Verdict:')],
    [TXT('        erro = max(abs(o - d) '), KW('for '), TXT('d, o, _ '), KW('in '),
     TXT('trace.pairs('), STR('"dac:1"'), TXT(', '), STR('"adc:1"'), TXT('))')],
    [TXT('        m = Measurement('), STR('"adc_erro_v"'), TXT(', erro, '), STR('"V"'),
     TXT(', limit_max='), NUM('0.05'), TXT(')')],
    [KW('        return '), TXT('Verdict(erro < '), NUM('0.05'), TXT(', '),
     STR('f"erro máx {erro:.3f} V"'), TXT(', [m])')],
], caption='Esqueleto de componentes novos (ex.: DAC → ADC em loopback)')
table(
    ['Regra', 'Por quê'],
    [
        ['name único por componente (ex.: "out:led", "in:botao")', 'O estímulo e as decisões '
         'referenciam drivers e monitores pelo nome'],
        ['setup() salva, teardown() restaura', 'O framework garante a chamada, inclusive com erro'],
        ['Estímulo sem acesso a hardware', 'A mesma sequência serve a qualquer driver'],
        ['Decisão devolve métrica com limite padrão', 'O perfil de critérios pode trocar o limite '
         'sem mudar o código (seção 16)'],
    ],
    widths=[7.0, 9.0],
    titulo='Regras para componentes do testbench',
)

h2('15.5 Critério de aceitação nas decisões')
para(
    'FollowsDrive e Expect registram a contagem de divergências como métrica, com limite padrão '
    '0: gpio_divergencias_<sinal>, gpio_nivel_divergente_<sinal>, gpio_pull_divergente_<sinal> e '
    'gpio_loopback_divergente_<...>. O veredito final vem dessas métricas: se o perfil de '
    'critérios do projeto definir, por exemplo, "gpio_divergencias" com limit_max = 1, uma '
    'divergência passa a ser aceita em todos os pinos — e o relatório registra a origem do limite.',
    space_after=6
)
box(
    'O testbench foi validado na NUCLEO-H753ZI (Piscar LED, GPIO Output/Input, Nível, Pull) e em\n'
    'testes unitários com um simulador de GPIO sobre o SVD real (curto a GND, fio de loopback,\n'
    'ativo-baixo). Critérios de aceitação de produto devem ser revisados e aprovados por pessoa\n'
    'distinta de quem escreveu o teste antes de virarem registro de qualidade.',
    fill='E8F4F8',
    title='  Validação',
    title_color=AZUL_MEDIO
)

quebra()


# ══════════════════════════════════════════════════════════════════════════════
# 16. TOLERÂNCIAS, RELATÓRIOS E LOGS
# ══════════════════════════════════════════════════════════════════════════════
h1('16. Tolerâncias, Relatórios e Logs')

h2('16.1 Perfil de critérios: o perfil vence')
para(
    'O perfil de critérios do projeto (botão "Critérios de análise...", arquivos em '
    'analysis_configs/) é aplicado a toda medida que entra num MeasurementSet. O limite escrito '
    'no teste vale como padrão; se o perfil define a métrica, o perfil vence.',
    space_after=6
)
table(
    ['Situação', 'Limite aplicado', 'Origem registrada'],
    [
        ['Perfil define "uart_ber" e a medida é "uart_ber_115200"', 'Do perfil (casa por prefixo)',
         'perfil Padrão Geral v1.0'],
        ['Perfil define tolerance_pct sem nominal', '± % sobre o nominal da medida', 'perfil ...'],
        ['Perfil não define a métrica', 'O do teste', 'padrão do teste'],
        ['Nenhum perfil vinculado ao projeto', 'O do teste', 'padrão do teste'],
    ],
    widths=[6.5, 5.0, 4.5],
    titulo='Como o limite de cada medida é escolhido',
)
para(
    'Cada medida guarda o valor nominal, a origem do limite e a margem: quanto ainda sobra até '
    'o limite mais próximo, em % da faixa (negativa = fora do limite). Na janela principal, '
    'passar o mouse sobre um resultado mostra esses dados.',
    space_after=8
)

h2('16.2 Relatórios')
bullet('"Exportar PDF" gera o relatório A4 direto (motor do Qt, sem dependência extra)')
bullet('"Relatório ▾" gera HTML, Word ou PDF; "Exportar..." gera JSON ou CSV')
bullet('A janela de salvar abre na última pasta usada; o nome leva projeto, data e hora')
bullet('Seções: resumo, resultados (com os parâmetros de cada teste), métricas com coluna '
       '"Margem", "Critérios de análise" (perfil ativo e suas métricas) e "Instrumentos '
       'utilizados"')

h2('16.3 Janela de logs')
para(
    'Menu Exibir → Logs (Ctrl+L), ou o contador "Logs" na barra de status, que fica vermelho '
    'quando há erros. A janela recebe tudo o que a plataforma registra, as exceções não '
    'tratadas (inclusive de threads) e os testes que terminam em FAIL ou ERROR, com os '
    'parâmetros usados.',
    space_after=6
)
figura('gui_logs.png',
       'Janela de logs: (1) filtro por nível e busca; (2) registros; (3) detalhe com traceback; '
       '(4) salvar e limpar (exemplo)')
table(
    ['Ação', 'Efeito'],
    [
        ['Salvar... (.jsonl)', 'JSON Lines: 1ª linha com o contexto da sessão (versão, projeto, '
                               'placa, probe, perfil); depois um registro por linha, com campos '
                               'separados — formato para análise automática'],
        ['Salvar... (.log)', 'Texto legível, com o contexto no cabeçalho'],
        ['Copiar', 'Registros filtrados para a área de transferência'],
        ['Limpar', 'Esvazia os logs da sessão; o arquivo de log em disco (modo dev) fica'],
    ],
    widths=[3.5, 12.5],
    titulo='Ações da janela de logs',
)
box(
    'Na janela e nos arquivos exportados, caminhos com a pasta do usuário (que contém o nome\n'
    'da pessoa) aparecem como "~" (LGPD). Para análise por IA, use somente as ferramentas\n'
    'aprovadas pelo Instituto (Claude Enterprise).',
    fill='E8F4F8',
    title='  Privacidade',
    title_color=AZUL_MEDIO
)

quebra()


# ══════════════════════════════════════════════════════════════════════════════
# 17. ERROS COMUNS
# ══════════════════════════════════════════════════════════════════════════════
h1('17. Erros Comuns e Como Resolver')

table(
    ['Sintoma', 'Causa Provável', 'Solução'],
    [
        [
            'RuntimeError: Nenhuma placa configurada',
            'get_active_board() chamado sem selecionar placa na GUI',
            'Selecione uma placa no ComboBox "Projeto" antes de conectar; '
            'em testes unitários, use set_active_board(BOARD) via fixture autouse'
        ],
        [
            'KeyError: Campo X',
            'Nome de periférico, registrador ou campo incorreto',
            'Consulte o SVD com o comando Python da seção 6.1'
        ],
        [
            'read_field() sempre retorna 0',
            'Clock do periférico não habilitado',
            'Adicione pal.modify_field("RCC", "AHBxENR", "XXXEN", 1) antes de acessar'
        ],
        [
            'Timeout no receive_byte()',
            'Fio de loopback não conectado ou pino não como AF',
            'Verifique a conexão física e confirme MODER=10 e AF correto'
        ],
        [
            'Teste passa mas o próximo falha',
            'O teste anterior não restaurou o estado dos pinos',
            'Use try/finally e salve/restaure todos os campos que modifica'
        ],
        [
            'Teste não aparece após "Carregar Plugin"',
            'Classe não herda de BuiltinTest ou é abstrata',
            'Verifique: class MeuTeste(BuiltinTest) e que _run() está implementado'
        ],
        [
            'Falha ao identificar MCU: Nenhum ST-Link encontrado',
            'Probe desconectado ou em uso por outro programa (STM32CubeIDE, CubeProgrammer)',
            'Feche o outro programa, reconecte o USB e clique "Atualizar"'
        ],
        [
            'Aviso "Não confere com a placa do projeto"',
            'O chip na bancada é diferente do BoardConfig do projeto selecionado',
            'Selecione (ou crie) o projeto da placa correta; confira o DEV_ID no manual de referência'
        ],
        [
            'Importar .ioc não seleciona a placa base',
            'Não há BoardConfig com target igual ao MCU do .ioc',
            'Crie src/boards/<placa>.py com target = Mcu.UserName em minúsculas'
        ],
        [
            'pyOCD: No cores were discovered',
            'Firmware em baixo consumo ou com o debug desligado',
            'Conecte com StLinkADIv5(..., mode="under-reset")'
        ],
        [
            'RuntimeError: SYSCLK não está no HSI',
            'O firmware já configurou os clocks (o chip não foi resetado)',
            'Chame t.reset_and_halt() antes de run_sweep()'
        ],
        [
            'RuntimeError: stub não iniciou (MAGIC ausente)',
            '.bin desatualizado, endereço de carga errado ou núcleo não liberado',
            'Rode python stubs/build_stubs.py e confira CODE_ADDR e o stub.ld'
        ],
        [
            'Leitura de bloco falha (TransferFaultError) a 24 MHz',
            'Clock SWD alto demais para o cabo/placa',
            'Use até 8 MHz de SWD'
        ],
        [
            'Teste composto em FAIL: "nenhuma amostra"',
            'O estímulo não amostrou o monitor (Step com sample=False) ou nome errado',
            'Confira os nomes "out:<sinal>", "in:<sinal>", "odr:<sinal>" usados nas decisões'
        ],
        [
            'Teste do plano em ERROR: parâmetros inválidos',
            'Sinal renomeado/removido no Pinout ou valor fora dos limites',
            'Editar plano de testes...: o formulário marca em vermelho o sinal que não existe'
        ],
        [
            'Bloco não lista o periférico (ex.: USART3)',
            'Pinos do periférico sem tipo/nome/função no Pinout',
            'Na aba Pinout, preencha Tipo (UART...), Periférico (USART3) e Função (TX/RX)'
        ],
        [
            'CAN em SKIP: firmware usando o CAN',
            'O controlador está fora do modo de inicialização',
            'Esperado: o bloco não derruba a comunicação do produto. Teste com firmware parado'
        ],
        [
            'Conexão H7: "Invalid coresight component"',
            'Descoberta do chip logo após a sessão anterior',
            'A plataforma repete até 3 vezes; se persistir, reconecte o USB'
        ],
        [
            'Medida sem o limite do perfil',
            'Nome da métrica não casa com o perfil',
            'Use o nome do perfil como prefixo: "uart_ber" vale para "uart_ber_115200"'
        ],
        [
            '"Detectar VISA..." não encontra o aparelho',
            'Aparelho em outra sub-rede, LAN desligada ou só USB sem biblioteca VISA',
            'Confira IP e máscara no aparelho; prefira LAN; pela USB instale o NI-VISA'
        ],
        [
            'UnsupportedModelError ao chamar get_oscilloscope()',
            'O *IDN? não corresponde a nenhum modelo em instruments.models.MODELS',
            'Defina params["model"] no cadastro ou inclua o modelo (seção 13.5)'
        ],
        [
            'CalibrationExpiredError ao conectar',
            'Calibração vencida com block_if_expired ativo',
            'Encaminhe o aparelho para calibração e atualize calibration_due'
        ],
        [
            'Teste de instrumento sempre SKIP',
            'Instrumento cadastrado com outro nome ou não conectado',
            'O nome no teste (ex.: "osc") deve ser igual ao do gerenciador; clique "Conectar"'
        ],
        [
            'Erro ao gerar o relatório Word',
            'python-docx não instalado no ambiente',
            'pip install python-docx (o PDF não depende de pacote extra)'
        ],
    ],
    widths=[4.5, 4.5, 7],
    titulo='Erros comuns e como resolver',
)

quebra()


# ══════════════════════════════════════════════════════════════════════════════
# 18. CHECKLIST
# ══════════════════════════════════════════════════════════════════════════════
h1('18. Checklist Final — Antes de Executar no Hardware')

para('Antes de conectar a placa e rodar o teste, confirme cada item:', space_after=10)

items = [
    ('Estrutura básica', [
        'A classe herda de BuiltinTest',
        'Os atributos "name" e "description" estão definidos',
        'O método "_run(self, pal)" está implementado (com underscore)',
        'Todo caminho de execução retorna self._pass(), self._fail() ou self._skip()',
    ]),
    ('Boas práticas', [
        'O estado original dos registradores é salvo antes de qualquer alteração',
        'O bloco finally restaura os registradores modificados',
        'O clock do periférico é habilitado via RCC antes de usar',
        'Todos os loops de espera têm timeout com time.perf_counter()',
    ]),
    ('Portabilidade', [
        'Pinos/clocks lidos de get_active_board() em testes de hardware (Padrão A)',
        'Constantes de firmware definidas no topo do arquivo (Padrão B)',
        'Teste não usa endereços/nomes hardcoded que mudariam ao trocar de placa',
        'Chip conferido com "Identificar MCU" (DEV_ID confere com a placa do projeto)',
        'Pinout importado do .ioc revisado (nomes de sinal e coluna "Ativo baixo")',
    ]),
    ('Plano de testes e blocos (seção 14)', [
        'Cada bloco tem name, description, category e params; a lógica não cita portas nem pinos',
        'Pinos de periférico com Tipo, Periférico e Função preenchidos no Pinout',
        'Plano salvo sem erros de parâmetro; entradas que exigem fio/jumper desmarcadas se ausentes',
        'Bloco restaura registradores e pinos que altera (hal.common.snapshot_pin / restore_pin)',
    ]),
    ('Testbench (seção 15)', [
        'build() só monta a bancada: sem acesso a hardware fora de drivers e monitores',
        'Nomes únicos de drivers/monitores; estímulo referencia esses nomes',
        'Todo driver/monitor restaura o pino em teardown()',
        'Decisões registram métrica com limite padrão (critério ajustável pelo perfil)',
    ]),
    ('Tolerâncias e logs (seção 16)', [
        'Perfil de critérios vinculado ao projeto; nomes das métricas casam com o perfil',
        'Relatório conferido: coluna Margem e seção Critérios de análise',
        'Logs da sessão salvos (.jsonl) junto ao relatório quando houver FAIL/ERROR',
    ]),
    ('Testes com instrumentos (seção 13)', [
        'O teste usa get_oscilloscope()/get_spectrum_analyzer() e faz _skip() se for None',
        'Patrimônio e validade de calibração preenchidos no cadastro do instrumento',
        'Aparelho conectado antes de executar (identificação vai para o relatório)',
        'Comandos do modelo conferidos com o manual e com o aparelho (validated=True)',
    ]),
    ('Testes com agente na RAM (seção 12)', [
        'O .bin foi regenerado com stubs/build_stubs.py após qualquer mudança no .c',
        'O layout da caixa de correio é idêntico no .c e no módulo Python',
        'O lado host foi testado com alvo simulado antes da placa',
        'O firmware volta a rodar no fim (reset_system) e o sha256 do stub foi registrado',
    ]),
    ('Qualidade', [
        'O teste foi testado com FakeMemory (sem hardware) primeiro',
        'Fixture autouse chama set_active_board(BOARD) nos testes unitários',
        'Os parâmetros "measured" e "expected" estão preenchidos no _pass/_fail',
    ]),
]

for titulo, subitems in items:
    para(titulo, bold=True, color=AZUL_MEDIO, space_before=8, space_after=2)
    for item in subitems:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item).font.size = Pt(11)

doc.add_paragraph()

box(
    'Arquivo de template pronto para copiar: DOC/TEMPLATE_TESTE.py\n'
    'Exemplos funcionais: TESTES_PLATAFORMA/ (IO, UART, SPI, I2C, ADC, etc.)\n'
    'Arquivo de configuração de placa: src/boards/nucleo_f446re.py\n'
    'Dúvidas ou sugestões: rafael.macedoengenharia@gmail.com',
    fill='E8F4F8',
    title='  Recursos',
    title_color=AZUL_MEDIO
)


# ══════════════════════════════════════════════════════════════════════════════
# Salvar
# ══════════════════════════════════════════════════════════════════════════════
out = 'DOC/GUIA_CRIACAO_DE_TESTES.docx'
doc.save(out)
print(f'Criado: {out}')

# Numeração das legendas, sumário e listas são campos: o Word calcula as páginas.
_PS = r'''
$ErrorActionPreference = "Stop"
$w = New-Object -ComObject Word.Application
$w.Visible = $false
try {
    $d = $w.Documents.Open($args[0], $false, $false)
    # Os números das legendas (SEQ) já saem certos do gerador; só sumário e listas.
    # Atualizar uma lista pode apagar marcadores (_Toc) da outra: repete até não
    # sobrar referência quebrada ("Erro! Indicador..." / "Error! Bookmark...").
    $ok = $false
    for ($i = 0; $i -lt 3 -and -not $ok; $i++) {
        foreach ($t in $d.TablesOfContents) { $t.Update() | Out-Null }
        foreach ($t in $d.TablesOfFigures) { $t.Update() | Out-Null }
        $d.Repaginate()
        foreach ($t in $d.TablesOfContents) { $t.UpdatePageNumbers() | Out-Null }
        foreach ($t in $d.TablesOfFigures) { $t.UpdatePageNumbers() | Out-Null }
        $txt = $d.Content.Text
        $ok = -not ($txt -match "Indicador n.o definido|Bookmark not defined")
    }
    $d.Save()
    if (-not $ok) { Write-Output "AVISO: referências quebradas no sumário/listas" }
    $d.Close($false)
} finally {
    $w.Quit()
    [void][Runtime.InteropServices.Marshal]::ReleaseComObject($w)
}
'''
if sys.platform == 'win32':
    import tempfile
    with tempfile.TemporaryDirectory() as _tmp:
        _ps1 = Path(_tmp) / 'atualizar_campos.ps1'
        _ps1.write_text(_PS, encoding='utf-8-sig')
        _r = subprocess.run(['powershell', '-NoProfile', '-ExecutionPolicy', 'Bypass',
                             '-File', str(_ps1), str(Path(out).resolve())],
                            capture_output=True, text=True, timeout=300)
    if _r.returncode == 0 and 'AVISO' in _r.stdout:
        print('Sumário e listas atualizados, mas com referências quebradas — '
              'no Word, Ctrl+A e F9 para corrigir')
    elif _r.returncode == 0:
        print('Sumário, listas e numeração atualizados pelo Word')
    else:
        print('Word indisponível — no Word, selecione tudo (Ctrl+A) e pressione F9 para '
              'atualizar sumário e listas.\n' + (_r.stderr or _r.stdout)[-400:])
else:
    print('Fora do Windows: no Word, selecione tudo (Ctrl+A) e pressione F9.')
